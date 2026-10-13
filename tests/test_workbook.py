"""Struktur- und Importtests für die Baucontrolling-Mappe."""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "Baucontrolling_Kaufmann.xlsx"

REQUIRED_SHEETS = {
    "24_Dashboard",
    "00_Navi",
    "01_Stammdaten",
    "05_Projekte",
    "13_Ausgangsrechnungen",
    "15_Zahlungen",
    "16_OPOS",
    "19_Cashflow",
    "20_Liquiditaet",
    "21_Forecast",
    "22_BWA",
    "23_SollIst",
    "26_Form_Rechnung",
    "30_Import",
}

REQUIRED_NAMES = {
    "FirmaName",
    "UStSatz",
    "MindestLiq",
    "AnfangsBank",
    "ProjektNr",
    "ARNr",
    "KundenNr",
}


def test_workbook_exists() -> None:
    assert XLSX.exists(), "Workbook fehlt — zuerst python3 -m baucontrolling.generate"


def test_sheets_and_names() -> None:
    wb = load_workbook(XLSX, read_only=True, data_only=False)
    missing = REQUIRED_SHEETS - set(wb.sheetnames)
    assert not missing, missing
    names = set(wb.defined_names.keys())
    missing_n = REQUIRED_NAMES - names
    assert not missing_n, missing_n
    wb.close()


def test_formulas_present() -> None:
    wb = load_workbook(XLSX, data_only=False)
    assert str(wb["24_Dashboard"].cell(5, 1).value).startswith("=")
    assert "SUMIFS" in str(wb["05_Projekte"]["U4"].value)
    assert "VLOOKUP" in str(wb["26_Form_Rechnung"]["B12"].value)
    assert "MindestLiq" in str(wb["20_Liquiditaet"]["B15"].value)
    wb.close()


def test_bank_import(tmp_path: Path | None = None) -> None:
    target = ROOT / "Baucontrolling_Kaufmann.xlsx"
    copy = Path("/tmp/baucontrolling_import_test.xlsx")
    copy.write_bytes(target.read_bytes())
    before = load_workbook(copy)
    used = sum(1 for r in range(4, 200) if before["15_Zahlungen"].cell(r, 1).value)
    before.close()
    proc = subprocess.run(
        [
            sys.executable,
            str(ROOT / "scripts/import_csv.py"),
            "--workbook",
            str(copy),
            "--typ",
            "bank",
            "--datei",
            str(ROOT / "templates/bankumsatz.csv"),
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    assert "3 Zeilen" in proc.stdout
    after = load_workbook(copy)
    ws = after["15_Zahlungen"]
    ids = [ws.cell(r, 1).value for r in range(4, 200) if ws.cell(r, 1).value]
    assert len(ids) == used + 3
    assert ids[-3:] == ["Z-017", "Z-018", "Z-019"]
    assert ws.cell(used + 4, 4).value == "AR-26004"
    after.close()
    copy.unlink(missing_ok=True)


if __name__ == "__main__":
    test_workbook_exists()
    test_sheets_and_names()
    test_formulas_present()
    test_bank_import()
    print("alle Tests ok")
