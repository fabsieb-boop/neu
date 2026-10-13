"""Navigation, Stammdaten und Stammdatenlisten."""

from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from baucontrolling import sample_data as sd
from baucontrolling.layout import (
    MAX_ROWS,
    add_table,
    dropdown,
    fill_blank_rows,
    named,
    new_sheet,
    write_headers,
    write_row,
)
from baucontrolling.styles import (
    ALIGN_L,
    FILL_CREAM,
    FILL_NAVY,
    FILL_STEEL,
    FILL_WHITE,
    FONT_BODY,
    FONT_H2,
    FONT_SMALL,
    FONT_WHITE,
    THIN,
    style_input,
)

SHEETS = [
    ("00_Navi", "Start und Blattübersicht"),
    ("01_Stammdaten", "Firma, Steuern, Parameter"),
    ("02_Kunden", "Auftraggeber"),
    ("03_Lieferanten", "Lieferanten und Nachunternehmer"),
    ("04_Personal", "Mitarbeiter und Stundensätze"),
    ("05_Projekte", "Bauvorhaben und Deckungsbeitrag"),
    ("06_LV", "Leistungsverzeichnis"),
    ("07_Nachtraege", "Nachträge / Claim"),
    ("08_Angebote", "Angebotsverfolgung"),
    ("09_Auftraege", "Beauftragte Verträge"),
    ("10_Stunden", "Zeiterfassung"),
    ("11_Material", "Material und Baustoffe"),
    ("12_Geraete", "Geräte und Fuhrpark"),
    ("13_Ausgangsrechnungen", "Abschlag, Schluss, Anzahlung"),
    ("14_Eingangsrechnungen", "Kreditoren"),
    ("15_Zahlungen", "Bank- und Kassenbewegungen"),
    ("16_OPOS", "Offene Posten Debitor / Kreditor"),
    ("17_Mahnwesen", "Mahnstufen und Fristen"),
    ("18_Gewaehrleistung", "Einbehalte und Avale"),
    ("19_Cashflow", "Cashflow nach Herkunft"),
    ("20_Liquiditaet", "13-Wochen- und 12-Monatsplan"),
    ("21_Forecast", "Umsatz- und Ergebnisforecast"),
    ("22_BWA", "Kurz-BWA / GuV"),
    ("23_SollIst", "Projekt-Soll-Ist"),
    ("24_Dashboard", "KPI-Cockpit"),
    ("25_Form_Angebot", "Formular Angebot"),
    ("26_Form_Rechnung", "Formular Rechnung"),
    ("27_Form_Abnahme", "Formular Abnahme"),
    ("28_Form_Bautage", "Bautagesbericht"),
    ("29_Form_Mahnung", "Formular Mahnung"),
    ("30_Import", "CSV- und DATEV-Import"),
    ("31_Anleitung", "Bedienung für die Buchhaltung"),
]


def add_navi(wb: Workbook) -> None:
    ws = new_sheet(
        wb,
        "00_Navi",
        "Bergmann Bau · Kaufmannsmappe",
        "Alle kaufmännischen Steuerungsblätter für das Baugewerbe · Geschäftsjahr 2026",
        8,
        {1: 28, 2: 42, 3: 18, 4: 18, 5: 18, 6: 18, 7: 18, 8: 18},
    )
    ws["A4"] = "So arbeiten Sie mit der Mappe"
    ws["A4"].font = FONT_H2
    hints = [
        "1. Gelbe bzw. weiße Zellen sind Eingaben. Blaue Zellen sind Formeln — nicht überschreiben.",
        "2. Neue Zeilen immer innerhalb der Tabelle einfügen (Tab-Taste am Tabellenende), dann bleiben SUMIFS erhalten.",
        "3. Formulare: oben die Beleg- oder Projektnummer wählen, der Druckbereich füllt sich automatisch.",
        "4. Importe: Blatt 30_Import und Skript scripts/import_csv.py. DATEV-, Bank- und Stunden-CSV möglich.",
        "5. Dashboard, Liquidität, Cashflow, Forecast und BWA rechnen aus den Bewegungsdaten — keine Doppelpflege.",
    ]
    for i, text in enumerate(hints):
        ws.merge_cells(start_row=5 + i, start_column=1, end_row=5 + i, end_column=6)
        cell = ws.cell(5 + i, 1, text)
        cell.font = FONT_BODY
        cell.alignment = ALIGN_L

    ws["A11"] = "Blatt"
    ws["B11"] = "Inhalt"
    for col in range(1, 3):
        ws.cell(11, col).font = FONT_WHITE
        ws.cell(11, col).fill = FILL_NAVY
        ws.cell(11, col).border = THIN

    for i, (name, desc) in enumerate(SHEETS):
        r = 12 + i
        cell = ws.cell(r, 1, name)
        cell.hyperlink = f"#{name}!A1"
        cell.font = Font(name="Calibri", size=10, color="1B365D", underline="single")
        ws.cell(r, 2, desc).font = FONT_BODY
        for c in range(1, 3):
            ws.cell(r, c).fill = FILL_WHITE if i % 2 == 0 else FILL_CREAM
            ws.cell(r, c).border = THIN

    ws["D11"] = "Kennzahlen-Sprung"
    ws["D11"].font = FONT_WHITE
    ws["D11"].fill = FILL_STEEL
    jumps = [
        ("KPI-Dashboard", "24_Dashboard"),
        ("Liquidität 13 Wochen", "20_Liquiditaet"),
        ("Cashflow", "19_Cashflow"),
        ("Forecast / Plan", "21_Forecast"),
        ("Offene Posten", "16_OPOS"),
        ("Projekt-Soll-Ist", "23_SollIst"),
    ]
    for i, (label, target) in enumerate(jumps):
        cell = ws.cell(12 + i, 4, label)
        cell.hyperlink = f"#{target}!A1"
        cell.font = Font(name="Calibri", size=10, color="FFFFFF", underline="single")
        cell.fill = FILL_STEEL
        cell.border = THIN

    ws.freeze_panes = "A12"


def add_stammdaten(wb: Workbook) -> None:
    ws = new_sheet(
        wb,
        "01_Stammdaten",
        "Stammdaten und Steuerungsparameter",
        "Einmal pflegen — alle Formulare, USt-Berechnungen und Liquiditätsgrenzen greifen hierauf zu",
        8,
        {1: 28, 2: 16, 3: 42, 4: 18, 5: 22, 6: 18, 7: 22, 8: 16},
    )
    labels = [
        (5, "Firmenname", FIRMA_VAL("name"), "FirmaName"),
        (6, "Straße", FIRMA_VAL("strasse"), "FirmaStrasse"),
        (7, "PLZ", FIRMA_VAL("plz"), "FirmaPLZ"),
        (8, "Ort", FIRMA_VAL("ort"), "FirmaOrt"),
        (9, "Telefon", FIRMA_VAL("tel"), "FirmaTel"),
        (10, "E-Mail", FIRMA_VAL("mail"), "FirmaMail"),
        (11, "USt-IdNr.", FIRMA_VAL("ustid"), "FirmaUStId"),
        (12, "Steuernummer", FIRMA_VAL("steuernr"), "FirmaStNr"),
        (13, "Handelsregister", FIRMA_VAL("hrb"), "FirmaHRB"),
        (14, "Geschäftsführer", FIRMA_VAL("gf"), "FirmaGF"),
        (15, "IBAN", FIRMA_VAL("iban"), "FirmaIBAN"),
        (16, "BIC", FIRMA_VAL("bic"), "FirmaBIC"),
        (17, "Bank", FIRMA_VAL("bank"), "FirmaBank"),
    ]
    ws["A4"] = "Firma"
    ws["A4"].font = FONT_H2
    for row, label, value, name in labels:
        ws.cell(row, 1, label).font = FONT_BODY
        ws.cell(row, 2, value)
        style_input(ws.cell(row, 2))
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        named(wb, name, f"'01_Stammdaten'!$B${row}")

    params = [
        (5, "Geschäftsjahr", sd.FIRMA["gj"], "int", "GJ"),
        (6, "USt-Satz", sd.FIRMA["ust"], "pct", "UStSatz"),
        (7, "Skonto %", sd.FIRMA["skonto_pct"], "pct", "SkontoPct"),
        (8, "Skonto Tage", sd.FIRMA["skonto_tage"], "int", "SkontoTage"),
        (9, "Zahlungsziel Kunden (Tage)", sd.FIRMA["zahlungsziel"], "int", "Zahlungsziel"),
        (10, "Sicherheitseinbehalt %", sd.FIRMA["einbehalt"], "pct", "EinbehaltPct"),
        (11, "Gewährleistung Jahre", sd.FIRMA["gewaehr_jahre"], "int", "GewaehrJahre"),
        (12, "Mindestliquidität", sd.FIRMA["mindestliquiditaet"], "eur", "MindestLiq"),
        (13, "Anfangsbestand Bank 01.01.", sd.FIRMA["anfangskasse"], "eur", "AnfangsBank"),
        (14, "Kreditlinie", sd.FIRMA["kreditlinie"], "eur", "Kreditlinie"),
        (15, "Sollstunden / Monat", sd.FIRMA["sollstunden_monat"], "qty", "SollStunden"),
        (16, "Mittl. Verrechnungssatz", sd.FIRMA["verrechnungssatz"], "eur", "VSSatz"),
        (17, "AGK-Zuschlag Angebot", sd.FIRMA["agk_pct"], "pct", "AGKPct"),
        (18, "Wagnis + Gewinn", sd.FIRMA["wagnis_pct"], "pct", "WuGPct"),
        (19, "Buchhaltung", sd.FIRMA["buchhaltung"], "text", "BuHaName"),
    ]
    ws["E4"] = "Parameter (wirken auf alle Blätter)"
    ws["E4"].font = FONT_H2
    for row, label, value, kind, name in params:
        ws.cell(row, 5, label).font = FONT_BODY
        ws.cell(row, 6, value)
        style_input(ws.cell(row, 6), kind)
        named(wb, name, f"'01_Stammdaten'!$F${row}")

    ws["A20"] = "Auswahl-Listen (Dropdown-Quellen)"
    ws["A20"].font = FONT_H2
    list_cols = [
        (1, "status_projekt", "lstProjektStatus"),
        (2, "status_beleg", "lstBelegStatus"),
        (3, "status_angebot", "lstAngebotStatus"),
        (4, "art_projekt", "lstProjektArt"),
        (5, "vertrag", "lstVertrag"),
        (6, "ar_typ", "lstARTyp"),
        (7, "er_typ", "lstERTyp"),
        (8, "zahlung_art", "lstZahlungArt"),
    ]
    for col, key, name in list_cols:
        ws.cell(21, col, key).font = FONT_WHITE
        ws.cell(21, col).fill = FILL_NAVY
        values = sd.LISTEN[key]
        for i, val in enumerate(values):
            ws.cell(22 + i, col, val)
            style_input(ws.cell(22 + i, col))
        last = 21 + len(values)
        letter = get_column_letter(col)
        named(wb, name, f"'01_Stammdaten'!${letter}$22:${letter}${last}")

    ws["A32"] = "Kostenarten (BWA / Soll-Ist)"
    ws["A32"].font = FONT_H2
    ka_headers = ["Nr", "Bezeichnung", "Gruppe", "SKR-Konto"]
    write_headers(ws, ka_headers, 33)
    for i, row in enumerate(sd.KOSTENARTEN):
        write_row(ws, 34 + i, row, ["text", "text", "text", "int"])
    add_table(ws, "tblKostenarten", f"A33:D{33 + len(sd.KOSTENARTEN)}")
    named(wb, "KostenartenGruppe", f"'01_Stammdaten'!$C$34:$C${33 + len(sd.KOSTENARTEN)}")
    ws.freeze_panes = "A4"


def FIRMA_VAL(key: str):
    return sd.FIRMA[key]


KUNDEN_H = ["Nr", "Name", "Typ", "Ansprechpartner", "Straße", "PLZ", "Ort", "Telefon", "E-Mail", "Zahlungsziel", "Skonto", "ABC", "Status", "Notiz"]
KUNDEN_K = ["text", "text", "text", "text", "text", "text", "text", "text", "text", "int", "pct", "text", "text", "text"]

LIEF_H = ["Nr", "Name", "Art", "Ansprechpartner", "IBAN", "Skonto Tage", "Skonto %", "Zahlungsziel", "E-Mail", "Status", "Notiz"]
LIEF_K = ["text", "text", "text", "text", "text", "int", "pct", "int", "text", "text", "text"]

PERS_H = ["Nr", "Name", "Funktion", "Kolonne", "Soll-h/Monat", "Verrechnung", "Kosten/h", "Eintritt", "Status", "Notiz"]
PERS_K = ["text", "text", "text", "text", "qty", "eur", "eur", "date", "text", "text"]


def add_kunden(wb: Workbook) -> None:
    ws = new_sheet(wb, "02_Kunden", "Kunden / Auftraggeber", "Debitorenstamm für Angebote, Aufträge, Rechnungen und Mahnungen", 14, _widths(KUNDEN_H))
    write_headers(ws, KUNDEN_H)
    for i, row in enumerate(sd.KUNDEN):
        write_row(ws, 4 + i, row, KUNDEN_K)
    fill_blank_rows(ws, 4 + len(sd.KUNDEN), MAX_ROWS, KUNDEN_K)
    add_table(ws, "tblKunden", f"A3:N{MAX_ROWS}")
    dropdown(ws, f"L4:L{MAX_ROWS}", "lstABC")
    dropdown(ws, f"M4:M{MAX_ROWS}", "lstBelegStatus")
    named(wb, "KundenNr", f"'02_Kunden'!$A$4:$A${MAX_ROWS}")
    named(wb, "KundenName", f"'02_Kunden'!$B$4:$B${MAX_ROWS}")
    ws.freeze_panes = "A4"


def add_lieferanten(wb: Workbook) -> None:
    ws = new_sheet(wb, "03_Lieferanten", "Lieferanten und Nachunternehmer", "Kreditorenstamm inkl. Skonto und Zahlungsziel für Liquiditätsplanung", 11, _widths(LIEF_H))
    write_headers(ws, LIEF_H)
    for i, row in enumerate(sd.LIEFERANTEN):
        write_row(ws, 4 + i, row, LIEF_K)
    fill_blank_rows(ws, 4 + len(sd.LIEFERANTEN), MAX_ROWS, LIEF_K)
    add_table(ws, "tblLieferanten", f"A3:K{MAX_ROWS}")
    named(wb, "LiefNr", f"'03_Lieferanten'!$A$4:$A${MAX_ROWS}")
    named(wb, "LiefName", f"'03_Lieferanten'!$B$4:$B${MAX_ROWS}")
    ws.freeze_panes = "A4"


def add_personal(wb: Workbook) -> None:
    ws = new_sheet(wb, "04_Personal", "Personal und Verrechnungssätze", "Kostensatz fließt in Stunden, Soll-Ist und BWA — Verrechnungssatz in Produktivitäts-KPI", 10, _widths(PERS_H))
    write_headers(ws, PERS_H)
    for i, row in enumerate(sd.PERSONAL):
        write_row(ws, 4 + i, row, PERS_K)
    fill_blank_rows(ws, 4 + len(sd.PERSONAL), MAX_ROWS, PERS_K)
    add_table(ws, "tblPersonal", f"A3:J{MAX_ROWS}")
    named(wb, "PersNr", f"'04_Personal'!$A$4:$A${MAX_ROWS}")
    named(wb, "PersName", f"'04_Personal'!$B$4:$B${MAX_ROWS}")
    ws.freeze_panes = "A4"
    ws["A2"] = "Mitarbeiterstamm · Kostensatz = Arbeitgebervollkosten je Stunde inkl. SV, Soka-Bau, BG"


def _widths(headers: list[str]) -> dict[int, float]:
    widths = {}
    for i, h in enumerate(headers, start=1):
        widths[i] = max(12, min(28, len(h) + 6))
    widths[1] = 14
    if len(headers) > 1:
        widths[2] = 28
    return widths


def add_all_stammdaten(wb: Workbook) -> None:
    add_navi(wb)
    add_stammdaten(wb)
    add_kunden(wb)
    add_lieferanten(wb)
    add_personal(wb)
