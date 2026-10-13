"""Druckbare kaufmännische Formulare mit Belegauswahl."""

from openpyxl.styles import Alignment, Font
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.page import PageMargins

from baucontrolling.layout import dropdown
from baucontrolling.styles import (
    ALIGN_L,
    FILL_CREAM,
    FILL_NAVY,
    FILL_STEEL,
    FILL_WHITE,
    FONT_BODY,
    FONT_H2,
    FONT_SMALL,
    FONT_WHITE,
    THIN,
    THIN_NAVY,
    paint_sheet,
    set_widths,
    style_calc,
    style_input,
)


def _form_base(wb: Workbook, title: str, headline: str) -> object:
    ws = wb.create_sheet(title)
    paint_sheet(ws, rows=60, cols=8)
    set_widths(ws, {1: 22, 2: 22, 3: 18, 4: 18, 5: 16, 6: 16, 7: 14, 8: 12})
    ws.merge_cells("A1:F1")
    ws["A1"] = headline
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = FILL_NAVY
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for col in range(1, 7):
        ws.cell(1, col).fill = FILL_NAVY
    ws.row_dimensions[1].height = 26
    ws.merge_cells("A2:F2")
    ws["A2"] = '=FirmaName&" · "&FirmaStrasse&" · "&FirmaPLZ&" "&FirmaOrt&" · USt-Id "&FirmaUStId'
    ws["A2"].font = FONT_SMALL
    ws["A2"].fill = FILL_STEEL
    ws["A2"].font = Font(name="Calibri", size=8, color="FFFFFF")
    for col in range(1, 7):
        ws.cell(2, col).fill = FILL_STEEL
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins = PageMargins(left=0.6, right=0.6, top=0.5, bottom=0.5)
    ws.print_area = "A1:F52"
    ws.sheet_view.showGridLines = False
    return ws


def _label(ws, row, col, text):
    cell = ws.cell(row, col, text)
    cell.font = FONT_SMALL
    return cell


def add_form_angebot(wb: Workbook) -> None:
    ws = _form_base(wb, "25_Form_Angebot", "ANGEBOT")
    ws["A4"] = "Angebots-Nr. wählen"
    style_input(ws["B4"])
    dropdown(ws, "B4", "AngebotNr")
    ws["B4"] = "AN-26007"
    labels = [
        (6, "Angebots-Nr.", "=IF(B4=\"\",\"\",B4)"),
        (7, "Datum", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'08_Angebote'!$A$4:$M$400,2,FALSE),\"\"))"),
        (8, "Kunde", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(VLOOKUP(B4,'08_Angebote'!$A$4:$M$400,3,FALSE),'02_Kunden'!$A$4:$B$400,2,FALSE),\"\"))"),
        (9, "Projekt", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'08_Angebote'!$A$4:$M$400,5,FALSE),\"\"))"),
        (10, "Netto Angebot", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'08_Angebote'!$A$4:$M$400,6,FALSE),0))"),
        (11, "AGK", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'08_Angebote'!$A$4:$M$400,7,FALSE),0))"),
        (12, "Wagnis + Gewinn", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'08_Angebote'!$A$4:$M$400,8,FALSE),0))"),
        (13, "Angebotssumme netto", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'08_Angebote'!$A$4:$M$400,12,FALSE),0))"),
        (14, "zzgl. USt", "=IF(B4=\"\",\"\",B13*UStSatz)"),
        (15, "Angebot brutto", "=IF(B4=\"\",\"\",B13+B14)"),
        (16, "Bindefrist bis", "=IF(B4=\"\",\"\",B7+IFERROR(VLOOKUP(B4,'08_Angebote'!$A$4:$M$400,11,FALSE),0))"),
        (17, "Status", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'08_Angebote'!$A$4:$M$400,9,FALSE),\"\"))"),
    ]
    for row, lab, formula in labels:
        _label(ws, row, 1, lab)
        kind = "eur" if row in (10, 13, 14, 15) else "pct" if row in (11, 12) else "date" if row in (7, 16) else "text"
        ws.cell(row, 2, formula)
        style_calc(ws.cell(row, 2), kind)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)

    ws["A19"] = "Leistungsbeschreibung (Kurztext)"
    ws["A19"].font = FONT_H2
    ws.merge_cells("A20:F24")
    ws["A20"] = "Wir bieten die im Leistungsverzeichnis beschriebenen Arbeiten gemäß VOB/B bzw. BGB zu umseitig genannten Bedingungen an. Nachträge nur nach schriftlicher Beauftragung."
    ws["A20"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A20"].fill = FILL_WHITE
    ws["A20"].border = THIN

    ws["A26"] = "Zahlungsbedingungen"
    ws["A27"] = '= "Abschläge nach Baufortschritt, Zahlungsziel "&Zahlungsziel&" Tage, Skonto "&TEXT(SkontoPct,"0.0%")&" bei Zahlung binnen "&SkontoTage&" Tagen. Sicherheitseinbehalt "&TEXT(EinbehaltPct,"0.0%")&"."'
    ws.merge_cells("A27:F28")
    ws["A27"].alignment = Alignment(wrap_text=True, vertical="top")

    ws["A30"] = "Ort, Datum"
    ws["C30"] = "Stempel / Unterschrift Bieter"
    ws["A31"] = '=FirmaOrt&", "&TEXT(TODAY(),"DD.MM.YYYY")'
    ws["A34"] = "Gültig ohne Unterschrift im Systemdruck — Original unterschrieben der Vergabe beilegen."
    ws["A34"].font = FONT_SMALL


def add_form_rechnung(wb: Workbook) -> None:
    ws = _form_base(wb, "26_Form_Rechnung", "RECHNUNG / ABSCHLAG")
    ws["A4"] = "Rechnungs-Nr. wählen"
    ws["B4"] = "AR-26004"
    style_input(ws["B4"])
    dropdown(ws, "B4", "ARNr")
    rows = [
        (6, "Rechnungs-Nr.", "=B4", "text"),
        (7, "Datum", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'13_Ausgangsrechnungen'!$A$4:$T$400,2,FALSE),\"\"))", "date"),
        (8, "Typ", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'13_Ausgangsrechnungen'!$A$4:$T$400,5,FALSE),\"\"))", "text"),
        (9, "Projekt", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'13_Ausgangsrechnungen'!$A$4:$T$400,3,FALSE),\"\"))", "text"),
        (10, "Projektname", "=IF(B9=\"\",\"\",IFERROR(VLOOKUP(B9,'05_Projekte'!$A$4:$B$400,2,FALSE),\"\"))", "text"),
        (11, "Kunde", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(VLOOKUP(B4,'13_Ausgangsrechnungen'!$A$4:$T$400,4,FALSE),'02_Kunden'!$A$4:$B$400,2,FALSE),\"\"))", "text"),
        (12, "Leistung netto", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'13_Ausgangsrechnungen'!$A$4:$T$400,7,FALSE),0))", "eur"),
        (13, "abzgl. Einbehalt", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'13_Ausgangsrechnungen'!$A$4:$T$400,14,FALSE),0))", "eur"),
        (14, "Netto fällig", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'13_Ausgangsrechnungen'!$A$4:$T$400,15,FALSE),0))", "eur"),
        (15, "USt", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'13_Ausgangsrechnungen'!$A$4:$T$400,16,FALSE),0))", "eur"),
        (16, "Brutto fällig", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'13_Ausgangsrechnungen'!$A$4:$T$400,17,FALSE),0))", "eur"),
        (17, "bereits bezahlt", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'13_Ausgangsrechnungen'!$A$4:$T$400,18,FALSE),0))", "eur"),
        (18, "Offen", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'13_Ausgangsrechnungen'!$A$4:$T$400,19,FALSE),0))", "eur"),
        (19, "Fällig am", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'13_Ausgangsrechnungen'!$A$4:$T$400,10,FALSE),\"\"))", "date"),
        (20, "Leistungsbeschreibung", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'13_Ausgangsrechnungen'!$A$4:$T$400,13,FALSE),\"\"))", "text"),
    ]
    for row, lab, formula, kind in rows:
        _label(ws, row, 1, lab)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        ws.cell(row, 2, formula)
        style_calc(ws.cell(row, 2), kind)
    ws["B16"].fill = FILL_CREAM
    ws["B18"].border = THIN_NAVY

    ws["A22"] = "Zahlungsempfänger"
    ws["A23"] = '=FirmaName&" · IBAN "&FirmaIBAN&" · BIC "&FirmaBIC&" · "&FirmaBank'
    ws.merge_cells("A23:F24")
    ws["A23"].alignment = Alignment(wrap_text=True)
    ws["A26"] = '= "Bitte überweisen Sie den offenen Betrag unter Angabe der Rechnungsnummer bis zum "&TEXT(B19,"DD.MM.YYYY")&". Bei Zahlung bis "&TEXT(B7+SkontoTage,"DD.MM.YYYY")&" gewähren wir "&TEXT(SkontoPct,"0.0%")&" Skonto."'
    ws.merge_cells("A26:F28")
    ws["A26"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A30"] = "Hinweis Einbehalt: Der Sicherheitseinbehalt wird bis zur Abnahme / Ablauf der Gewährleistung einbehalten bzw. gegen Bürgschaft ausgezahlt."
    ws.merge_cells("A30:F31")
    ws["A30"].font = FONT_SMALL


def add_form_abnahme(wb: Workbook) -> None:
    ws = _form_base(wb, "27_Form_Abnahme", "ABNAHMEPROTOKOLL")
    ws["A4"] = "Projekt wählen"
    ws["B4"] = "P-2026-005"
    style_input(ws["B4"])
    dropdown(ws, "B4", "ProjektNr")
    fields = [
        (6, "Projekt", "=B4"),
        (7, "Bezeichnung", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'05_Projekte'!$A$4:$T$400,2,FALSE),\"\"))"),
        (8, "Auftraggeber", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(VLOOKUP(B4,'05_Projekte'!$A$4:$C$400,3,FALSE),'02_Kunden'!$A$4:$B$400,2,FALSE),\"\"))"),
        (9, "Vertragsart", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'05_Projekte'!$A$4:$T$400,5,FALSE),\"\"))"),
        (10, "Bauleiter", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'05_Projekte'!$A$4:$R$400,17,FALSE),\"\"))"),
        (11, "Ort der Baustelle", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'05_Projekte'!$A$4:$T$400,19,FALSE),\"\"))"),
        (12, "Auftrag netto", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'05_Projekte'!$A$4:$V$400,22,FALSE),0))"),
    ]
    for row, lab, formula in fields:
        _label(ws, row, 1, lab)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        ws.cell(row, 2, formula)
        style_calc(ws.cell(row, 2), "eur" if row == 12 else "text")
    ws["A14"] = "Abnahmedatum"
    style_input(ws["B14"], "date")
    ws["A15"] = "Art der Abnahme"
    ws["B15"] = "förmlich nach VOB/B § 12"
    style_input(ws["B15"])
    ws["A17"] = "Feststellungen / Mängel"
    ws.merge_cells("A18:F24")
    ws["A18"] = "Keine wesentlichen Mängel. Restarbeiten gemäß Mängelliste Anlage 1 innerhalb von 14 Werktagen."
    ws["A18"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A18"].fill = FILL_WHITE
    ws["A18"].border = THIN
    ws["A26"] = "Beginn Gewährleistung"
    ws["B26"] = "=IF(B14=\"\",\"\",B14)"
    style_calc(ws["B26"], "date")
    ws["A27"] = "Ende Gewährleistung"
    ws["B27"] = "=IF(B14=\"\",\"\",DATE(YEAR(B14)+GewaehrJahre,MONTH(B14),DAY(B14)))"
    style_calc(ws["B27"], "date")
    ws["A29"] = "Auftragnehmer"
    ws["D29"] = "Auftraggeber"
    ws["A30"] = "=FirmaGF"
    style_calc(ws["A30"], "text")
    ws["A32"] = "______________________________"
    ws["D32"] = "______________________________"
    ws["A33"] = "Unterschrift / Datum"
    ws["D33"] = "Unterschrift / Datum"


def add_form_bautage(wb: Workbook) -> None:
    ws = _form_base(wb, "28_Form_Bautage", "BAUTAGESBERICHT")
    ws["A4"] = "Projekt"
    ws["B4"] = "P-2026-001"
    style_input(ws["B4"])
    dropdown(ws, "B4", "ProjektNr")
    ws["A5"] = "Datum"
    style_input(ws["B5"], "date")
    from datetime import date

    ws["B5"] = date(2026, 10, 12)
    ws["A6"] = "Bauleiter / Polier"
    ws["B6"] = "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'05_Projekte'!$A$4:$R$400,17,FALSE),\"\"))"
    style_calc(ws["B6"], "text")
    ws["C6"] = "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'05_Projekte'!$A$4:$R$400,18,FALSE),\"\"))"
    style_calc(ws["C6"], "text")
    ws["A8"] = "Wetter / Temperatur"
    style_input(ws["B8"])
    ws["B8"] = "heiter, 12 °C, trocken"
    ws["A9"] = "Personal auf der Baustelle"
    ws["B9"] = "=IF(B4=\"\",\"\",COUNTIFS('10_Stunden'!C:C,B4,'10_Stunden'!B:B,B5))"
    style_calc(ws["B9"], "qty")
    ws["A10"] = "Stunden gesamt"
    ws["B10"] = "=IF(B4=\"\",\"\",SUMIFS('10_Stunden'!E:E,'10_Stunden'!C:C,B4,'10_Stunden'!B:B,B5))"
    style_calc(ws["B10"], "hrs")
    ws["A12"] = "Ausgeführte Leistungen"
    ws.merge_cells("A13:F18")
    ws["A13"] = "Rohbau Block C: Mauerwerk OG1, Deckenbewehrung vorbereitet. Nachunternehmer Elektro in Steigzone."
    ws["A13"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A13"].fill = FILL_WHITE
    ws["A20"] = "Geräte / Material / Behinderungen"
    ws.merge_cells("A21:F24")
    ws["A21"] = "Kran 81K in Betrieb. Betonlieferung 08:30. Keine Behinderung."
    ws["A21"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A21"].fill = FILL_WHITE
    ws["A26"] = "Unfälle / besondere Vorkommnisse: keine"
    ws["A28"] = "Unterschrift Bauleitung"
    ws["D28"] = "Kenntnisnahme AG (optional)"


def add_form_mahnung(wb: Workbook) -> None:
    ws = _form_base(wb, "29_Form_Mahnung", "MAHNUNG")
    ws["A4"] = "Mahnung wählen"
    ws["B4"] = "MAH-02"
    style_input(ws["B4"])
    dropdown(ws, "B4", "MahnungNr")
    rows = [
        (6, "Mahnung", "=B4", "text"),
        (7, "Rechnung", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'17_Mahnwesen'!$A$4:$K$400,2,FALSE),\"\"))", "text"),
        (8, "Kunde", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(VLOOKUP(B4,'17_Mahnwesen'!$A$4:$C$400,3,FALSE),'02_Kunden'!$A$4:$B$400,2,FALSE),\"\"))", "text"),
        (9, "Mahnstufe", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'17_Mahnwesen'!$A$4:$K$400,7,FALSE),\"\"))", "text"),
        (10, "Offener Betrag", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'17_Mahnwesen'!$A$4:$K$400,11,FALSE),0))", "eur"),
        (11, "Neue Frist", "=IF(B4=\"\",\"\",IFERROR(VLOOKUP(B4,'17_Mahnwesen'!$A$4:$K$400,10,FALSE),\"\"))", "date"),
    ]
    for row, lab, formula, kind in rows:
        _label(ws, row, 1, lab)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        ws.cell(row, 2, formula)
        style_calc(ws.cell(row, 2), kind)
    ws.merge_cells("A13:F20")
    ws["A13"] = (
        '="Sehr geehrte Damen und Herren,"&CHAR(10)&CHAR(10)&'
        '"trotz fälliger Rechnung "&B7&" ist der Betrag von "&TEXT(B10,"#,##0.00 €")&" noch offen. "'
        '&"Wir fordern Sie auf, den Betrag bis spätestens "&TEXT(B11,"DD.MM.YYYY")&" auf unser Konto "&FirmaIBAN&" zu überweisen. "'
        '&"Bei weiterem Verzug behalten wir uns Verzugszinsen, Mahnkosten und die Abgabe an unseren Rechtsanwalt vor."&CHAR(10)&CHAR(10)&'
        '"Mit freundlichen Grüßen"&CHAR(10)&FirmaName&CHAR(10)&BuHaName'
    )
    ws["A13"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A13"].fill = FILL_WHITE
    ws.row_dimensions[13].height = 120


def add_all_formulare(wb: Workbook) -> None:
    add_form_angebot(wb)
    add_form_rechnung(wb)
    add_form_abnahme(wb)
    add_form_bautage(wb)
    add_form_mahnung(wb)
