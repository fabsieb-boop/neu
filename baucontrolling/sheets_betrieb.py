"""Stunden, Material, Geräte."""

from openpyxl.workbook.workbook import Workbook

from baucontrolling import sample_data as sd
from baucontrolling.layout import MAX_ROWS, add_table, dropdown, named, new_sheet, write_headers, write_row

STD_H = [
    "Nr", "Datum", "Projekt", "Mitarbeiter", "Stunden", "Gewerk", "Tätigkeit", "Verrechenbar",
    "Kostensatz", "Verrechnung", "Kosten", "Leistungswert", "Deckungsbeitrag Std.",
]
MAT_H = [
    "Nr", "Datum", "Projekt", "Lieferant", "ER-Nr", "Bezeichnung", "ME", "Menge", "EP netto",
    "Kategorie", "Status", "Summe netto",
]
GER_H = [
    "Nr", "Bezeichnung", "Typ", "Eigentum", "Anschaffung", "Inbetrieb", "ND Jahre",
    "Satz/h", "Einsatz", "Status", "AfA p.a.", "Restbuchwert 2026",
]


def add_stunden(wb: Workbook) -> None:
    ws = new_sheet(
        wb,
        "10_Stunden",
        "Zeiterfassung",
        "Kosten = Stunden × Kostensatz Mitarbeiter · Leistungswert = Stunden × Verrechnungssatz bei verrechenbar",
        13,
        {1: 12, 2: 12, 3: 14, 4: 14, 5: 10, 6: 14, 7: 28, 8: 12, 9: 12, 10: 13, 11: 12, 12: 14, 13: 16},
    )
    write_headers(ws, STD_H)
    kinds = ["text", "date", "text", "text", "hrs", "text", "text", "text", "calc:eur", "calc:eur", "calc:eur", "calc:eur", "calc:eur"]
    for i, src in enumerate(sd.STUNDEN):
        r = 4 + i
        verrechenbar = "JA" if src[7] else "NEIN"
        vals = list(src[:7]) + [verrechenbar] + _std_calc(r)
        write_row(ws, r, vals, kinds)
    for r in range(4 + len(sd.STUNDEN), MAX_ROWS + 1):
        write_row(ws, r, [""] * 8 + _std_calc(r), kinds)
    add_table(ws, "tblStunden", f"A3:M{MAX_ROWS}")
    dropdown(ws, f"C4:C{MAX_ROWS}", "ProjektNr")
    dropdown(ws, f"D4:D{MAX_ROWS}", "PersNr")
    dropdown(ws, f"H4:H{MAX_ROWS}", '"JA,NEIN"')
    named(wb, "StundenKosten", f"'10_Stunden'!$K$4:$K${MAX_ROWS}")
    ws.freeze_panes = "A4"


def _std_calc(r: int) -> list[str]:
    return [
        f'=IF(D{r}="","",IFERROR(VLOOKUP(D{r},\'04_Personal\'!$A$4:$G${MAX_ROWS},7,FALSE),0))',
        f'=IF(D{r}="","",IFERROR(VLOOKUP(D{r},\'04_Personal\'!$A$4:$G${MAX_ROWS},6,FALSE),VSSatz))',
        f'=IF(A{r}="","",E{r}*I{r})',
        f'=IF(A{r}="","",IF(H{r}="JA",E{r}*J{r},0))',
        f'=IF(A{r}="","",L{r}-K{r})',
    ]


def add_material(wb: Workbook) -> None:
    ws = new_sheet(
        wb,
        "11_Material",
        "Material und Baustoffeinsatz",
        "Summe netto = Menge × EP · Zuordnung zur Eingangsrechnung über ER-Nr für den Abgleich Einkauf / Rechnung",
        12,
        {1: 12, 2: 12, 3: 14, 4: 12, 5: 12, 6: 28, 7: 10, 8: 10, 9: 12, 10: 14, 11: 12, 12: 14},
    )
    write_headers(ws, MAT_H)
    kinds = ["text", "date", "text", "text", "text", "text", "text", "qty", "eur", "text", "text", "calc:eur"]
    for i, src in enumerate(sd.MATERIAL):
        r = 4 + i
        vals = list(src) + [f'=IF(A{r}="","",H{r}*I{r})']
        write_row(ws, r, vals, kinds)
    for r in range(4 + len(sd.MATERIAL), MAX_ROWS + 1):
        write_row(ws, r, [""] * 11 + [f'=IF(A{r}="","",H{r}*I{r})'], kinds)
    add_table(ws, "tblMaterial", f"A3:L{MAX_ROWS}")
    dropdown(ws, f"C4:C{MAX_ROWS}", "ProjektNr")
    dropdown(ws, f"D4:D{MAX_ROWS}", "LiefNr")
    ws.freeze_panes = "A4"


def add_geraete(wb: Workbook) -> None:
    ws = new_sheet(
        wb,
        "12_Geraete",
        "Geräte und Fuhrpark",
        "Lineare AfA, Restbuchwert und internem Verrechnungssatz für Kalkulation und BWA",
        12,
        {1: 10, 2: 28, 3: 12, 4: 12, 5: 14, 6: 12, 7: 10, 8: 12, 9: 14, 10: 14, 11: 12, 12: 16},
    )
    write_headers(ws, GER_H)
    kinds = ["text", "text", "text", "text", "eur", "date", "int", "eur", "text", "text", "calc:eur", "calc:eur"]
    for i, src in enumerate(sd.GERAETE):
        r = 4 + i
        vals = list(src) + [
            f'=IF(OR(A{r}="",E{r}="",G{r}=0),"",E{r}/G{r})',
            f'=IF(A{r}="","",MAX(0,E{r}-K{r}*MAX(0,GJ-YEAR(F{r}))))',
        ]
        write_row(ws, r, vals, kinds)
    for r in range(4 + len(sd.GERAETE), MAX_ROWS + 1):
        write_row(
            ws,
            r,
            [""] * 10
            + [
                f'=IF(OR(A{r}="",E{r}="",G{r}=0),"",E{r}/G{r})',
                f'=IF(A{r}="","",MAX(0,E{r}-K{r}*MAX(0,GJ-YEAR(F{r}))))',
            ],
            kinds,
        )
    add_table(ws, "tblGeraete", f"A3:L{MAX_ROWS}")
    dropdown(ws, f"I4:I{MAX_ROWS}", "ProjektNr")
    ws.freeze_panes = "A4"


def add_all_betrieb(wb: Workbook) -> None:
    add_stunden(wb)
    add_material(wb)
    add_geraete(wb)
