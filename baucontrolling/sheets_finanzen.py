"""Rechnungen, Zahlungen, OPOS, Mahnwesen, Gewährleistung."""

from openpyxl.workbook.workbook import Workbook

from baucontrolling import sample_data as sd
from baucontrolling.layout import MAX_ROWS, add_table, dropdown, named, new_sheet, status_colors, write_headers, write_row

AR_H = [
    "Nr", "Datum", "Projekt", "Kunde", "Typ", "Abschlag-Nr", "Netto", "Einbehalt %", "USt %",
    "Fällig", "Bezahlt am", "Status", "Text",
    "Einbehalt EUR", "Netto fällig", "USt EUR", "Brutto", "Bezahlt", "Offen", "Tage überfällig",
]
ER_H = [
    "Nr", "Datum", "Projekt", "Lieferant", "Typ", "Netto", "USt %", "Fällig", "Bezahlt am",
    "Status", "Text", "USt EUR", "Brutto", "Bezahlt", "Offen", "Tage überfällig",
]
ZAH_H = ["Nr", "Datum", "Art", "Beleg", "Partner", "Betrag", "Weg", "Text", "Monat", "Jahr"]
MAH_H = ["Nr", "Rechnung", "Kunde", "Datum", "Stufe", "Frist Tage", "Art", "Status", "Text", "Fällig Mahnung", "Rechnungs-Offen"]
GW_H = ["Nr", "Projekt", "Beleg", "Beginn", "Ende", "Betrag", "Art", "Status", "Text", "Restlaufzeit Jahre"]


def _ar_calc(r: int) -> list[str]:
    return [
        f'=IF(A{r}="","",G{r}*H{r})',
        f'=IF(A{r}="","",G{r}-N{r})',
        f'=IF(A{r}="","",O{r}*I{r})',
        f'=IF(A{r}="","",O{r}+P{r})',
        f'=IF(A{r}="","",SUMIFS(\'15_Zahlungen\'!$F:$F,\'15_Zahlungen\'!$D:$D,A{r},\'15_Zahlungen\'!$C:$C,"Eingang"))',
        f'=IF(A{r}="","",MAX(0,Q{r}-R{r}))',
        f'=IF(OR(A{r}="",L{r}="bezahlt",J{r}=""),"",MAX(0,TODAY()-J{r}))',
    ]


def add_ar(wb: Workbook) -> None:
    widths = {i: 13 for i in range(1, 21)}
    widths.update({1: 12, 13: 28, 3: 14})
    ws = new_sheet(
        wb,
        "13_Ausgangsrechnungen",
        "Ausgangsrechnungen (Abschlag, Anzahlung, Schluss)",
        "Einbehalt und USt rechnen automatisch · Bezahlt kommt aus dem Zahlungsblatt · Offen speist OPOS, Liquidität und Dashboard",
        20,
        widths,
    )
    write_headers(ws, AR_H)
    kinds = [
        "text", "date", "text", "text", "text", "int", "eur", "pct", "pct",
        "date", "date", "text", "text",
        "calc:eur", "calc:eur", "calc:eur", "calc:eur", "calc:eur", "calc:eur", "calc:qty",
    ]
    for i, src in enumerate(sd.AR):
        r = 4 + i
        write_row(ws, r, list(src) + _ar_calc(r), kinds)
    for r in range(4 + len(sd.AR), MAX_ROWS + 1):
        write_row(ws, r, [""] * 13 + _ar_calc(r), kinds)
    add_table(ws, "tblAR", f"A3:T{MAX_ROWS}")
    dropdown(ws, f"C4:C{MAX_ROWS}", "ProjektNr")
    dropdown(ws, f"D4:D{MAX_ROWS}", "KundenNr")
    dropdown(ws, f"E4:E{MAX_ROWS}", "lstARTyp")
    dropdown(ws, f"L4:L{MAX_ROWS}", "lstBelegStatus")
    status_colors(ws, 12, 4, MAX_ROWS)
    named(wb, "ARNr", f"'13_Ausgangsrechnungen'!$A$4:$A${MAX_ROWS}")
    ws.freeze_panes = "A4"


def _er_calc(r: int) -> list[str]:
    return [
        f'=IF(A{r}="","",F{r}*G{r})',
        f'=IF(A{r}="","",F{r}+L{r})',
        f'=IF(A{r}="","",SUMIFS(\'15_Zahlungen\'!$F:$F,\'15_Zahlungen\'!$D:$D,A{r},\'15_Zahlungen\'!$C:$C,"Ausgang"))',
        f'=IF(A{r}="","",MAX(0,M{r}-N{r}))',
        f'=IF(OR(A{r}="",J{r}="bezahlt",H{r}=""),"",MAX(0,TODAY()-H{r}))',
    ]


def add_er(wb: Workbook) -> None:
    ws = new_sheet(
        wb,
        "14_Eingangsrechnungen",
        "Eingangsrechnungen / Kreditoren",
        "Material, Nachunternehmer, Gerät — steuert Cash-out, BWA und offene Verbindlichkeiten",
        16,
        {1: 12, 2: 12, 3: 14, 4: 12, 5: 16, 6: 14, 7: 10, 8: 12, 9: 12, 10: 12, 11: 22, 12: 12, 13: 12, 14: 12, 15: 12, 16: 14},
    )
    write_headers(ws, ER_H)
    kinds = [
        "text", "date", "text", "text", "text", "eur", "pct", "date", "date",
        "text", "text", "calc:eur", "calc:eur", "calc:eur", "calc:eur", "calc:qty",
    ]
    for i, src in enumerate(sd.ER):
        r = 4 + i
        write_row(ws, r, list(src) + _er_calc(r), kinds)
    for r in range(4 + len(sd.ER), MAX_ROWS + 1):
        write_row(ws, r, [""] * 11 + _er_calc(r), kinds)
    add_table(ws, "tblER", f"A3:P{MAX_ROWS}")
    dropdown(ws, f"C4:C{MAX_ROWS}", "ProjektNr")
    dropdown(ws, f"D4:D{MAX_ROWS}", "LiefNr")
    dropdown(ws, f"E4:E{MAX_ROWS}", "lstERTyp")
    dropdown(ws, f"J4:J{MAX_ROWS}", "lstBelegStatus")
    status_colors(ws, 10, 4, MAX_ROWS)
    named(wb, "ERNr", f"'14_Eingangsrechnungen'!$A$4:$A${MAX_ROWS}")
    ws.freeze_panes = "A4"


def add_zahlungen(wb: Workbook) -> None:
    ws = new_sheet(
        wb,
        "15_Zahlungen",
        "Zahlungen und Bankumsätze",
        "Importziel für Bank-CSV · Belegnummer verknüpft automatisch Debitoren- und Kreditoren-OP",
        10,
        {1: 10, 2: 12, 3: 12, 4: 12, 5: 16, 6: 14, 7: 10, 8: 28, 9: 10, 10: 10},
    )
    write_headers(ws, ZAH_H)
    kinds = ["text", "date", "text", "text", "text", "eur", "text", "text", "calc:qty", "calc:qty"]
    for i, src in enumerate(sd.ZAHLUNGEN):
        r = 4 + i
        vals = list(src) + [f'=IF(B{r}="","",MONTH(B{r}))', f'=IF(B{r}="","",YEAR(B{r}))']
        write_row(ws, r, vals, kinds)
    for r in range(4 + len(sd.ZAHLUNGEN), MAX_ROWS + 1):
        write_row(ws, r, [""] * 8 + [f'=IF(B{r}="","",MONTH(B{r}))', f'=IF(B{r}="","",YEAR(B{r}))'], kinds)
    add_table(ws, "tblZahlungen", f"A3:J{MAX_ROWS}")
    dropdown(ws, f"C4:C{MAX_ROWS}", "lstZahlungArt")
    ws.freeze_panes = "A4"


def add_opos(wb: Workbook) -> None:
    ws = new_sheet(
        wb,
        "16_OPOS",
        "Offene Posten Debitoren und Kreditoren",
        "Lebt ausschließlich von Formeln — nichts manuell eintragen. Filter nach überfällig für den Jour fixe.",
        10,
        {1: 16, 2: 16, 3: 14, 4: 16, 5: 14, 6: 14, 7: 14, 8: 14, 9: 16, 10: 14},
    )
    write_headers(ws, ["Seite", "Beleg", "Datum", "Partner", "Projekt", "Typ", "Brutto", "Offen", "Fällig", "Tage"])
    # Debitoren-Spiegel der AR-Zeilen
    kinds = ["calc:text", "calc:text", "calc:date", "calc:text", "calc:text", "calc:text", "calc:eur", "calc:eur", "calc:date", "calc:qty"]
    for r in range(4, MAX_ROWS + 1):
        src = r
        vals = [
            f'=IF(\'13_Ausgangsrechnungen\'!A{src}="","",IF(\'13_Ausgangsrechnungen\'!S{src}>0.5,"Debitor",""))',
            f'=IF(A{r}="","",\'13_Ausgangsrechnungen\'!A{src})',
            f'=IF(A{r}="","",\'13_Ausgangsrechnungen\'!B{src})',
            f'=IF(A{r}="","",\'13_Ausgangsrechnungen\'!D{src})',
            f'=IF(A{r}="","",\'13_Ausgangsrechnungen\'!C{src})',
            f'=IF(A{r}="","",\'13_Ausgangsrechnungen\'!E{src})',
            f'=IF(A{r}="","",\'13_Ausgangsrechnungen\'!Q{src})',
            f'=IF(A{r}="","",\'13_Ausgangsrechnungen\'!S{src})',
            f'=IF(A{r}="","",\'13_Ausgangsrechnungen\'!J{src})',
            f'=IF(A{r}="","",\'13_Ausgangsrechnungen\'!T{src})',
        ]
        write_row(ws, r, vals, kinds)
    # Kreditoren daneben
    ws.cell(3, 12, "Seite")
    extra = ["Seite", "Beleg", "Datum", "Partner", "Projekt", "Typ", "Brutto", "Offen", "Fällig", "Tage"]
    for i, h in enumerate(extra, start=12):
        ws.cell(3, i, h)
    from baucontrolling.styles import style_header_row

    style_header_row(ws, 3, 10, start=12)
    for r in range(4, MAX_ROWS + 1):
        src = r
        kred = [
            f'=IF(\'14_Eingangsrechnungen\'!A{src}="","",IF(\'14_Eingangsrechnungen\'!O{src}>0.5,"Kreditor",""))',
            f'=IF(L{r}="","",\'14_Eingangsrechnungen\'!A{src})',
            f'=IF(L{r}="","",\'14_Eingangsrechnungen\'!B{src})',
            f'=IF(L{r}="","",\'14_Eingangsrechnungen\'!D{src})',
            f'=IF(L{r}="","",\'14_Eingangsrechnungen\'!C{src})',
            f'=IF(L{r}="","",\'14_Eingangsrechnungen\'!E{src})',
            f'=IF(L{r}="","",\'14_Eingangsrechnungen\'!M{src})',
            f'=IF(L{r}="","",\'14_Eingangsrechnungen\'!O{src})',
            f'=IF(L{r}="","",\'14_Eingangsrechnungen\'!H{src})',
            f'=IF(L{r}="","",\'14_Eingangsrechnungen\'!P{src})',
        ]
        kinds_k = ["calc:text", "calc:text", "calc:date", "calc:text", "calc:text", "calc:text", "calc:eur", "calc:eur", "calc:date", "calc:qty"]
        for c, (v, k) in enumerate(zip(kred, kinds_k), start=12):
            from baucontrolling.styles import style_calc

            cell = ws.cell(r, c, v)
            style_calc(cell, k.split(":")[1])
    ws["A2"] = "Links Debitoren · rechts Kreditoren · Zeilen ohne Offenbetrag bleiben leer"
    ws.freeze_panes = "A4"
    from openpyxl.utils import get_column_letter

    for col, w in {12: 12, 13: 12, 14: 12, 15: 12, 16: 14, 17: 14, 18: 12, 19: 12, 20: 12, 21: 12}.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def add_mahnwesen(wb: Workbook) -> None:
    ws = new_sheet(
        wb,
        "17_Mahnwesen",
        "Mahnwesen",
        "Stufe und Frist · offener Rechnungsbetrag wird gezogen · Formular 29_Form_Mahnung druckt die gewählte Mahnung",
        11,
        {1: 12, 2: 12, 3: 12, 4: 12, 5: 10, 6: 12, 7: 24, 8: 12, 9: 28, 10: 14, 11: 16},
    )
    write_headers(ws, MAH_H)
    kinds = ["text", "text", "text", "date", "int", "int", "text", "text", "text", "calc:date", "calc:eur"]
    for i, src in enumerate(sd.MAHNUNGEN):
        r = 4 + i
        vals = list(src) + [
            f'=IF(A{r}="","",D{r}+F{r})',
            f'=IF(A{r}="","",IFERROR(VLOOKUP(B{r},\'13_Ausgangsrechnungen\'!$A$4:$S${MAX_ROWS},19,FALSE),0))',
        ]
        write_row(ws, r, vals, kinds)
    for r in range(4 + len(sd.MAHNUNGEN), MAX_ROWS + 1):
        write_row(
            ws,
            r,
            [""] * 9
            + [
                f'=IF(A{r}="","",D{r}+F{r})',
                f'=IF(A{r}="","",IFERROR(VLOOKUP(B{r},\'13_Ausgangsrechnungen\'!$A$4:$S${MAX_ROWS},19,FALSE),0))',
            ],
            kinds,
        )
    add_table(ws, "tblMahnungen", f"A3:K{MAX_ROWS}")
    dropdown(ws, f"B4:B{MAX_ROWS}", "ARNr")
    dropdown(ws, f"C4:C{MAX_ROWS}", "KundenNr")
    dropdown(ws, f"G4:G{MAX_ROWS}", '"Zahlungserinnerung,1. Mahnung,2. Mahnung + Verzug,3. Mahnung / RA"')
    named(wb, "MahnungNr", f"'17_Mahnwesen'!$A$4:$A${MAX_ROWS}")
    ws.freeze_panes = "A4"


def add_gewaehr(wb: Workbook) -> None:
    ws = new_sheet(
        wb,
        "18_Gewaehrleistung",
        "Sicherheitseinbehalte und Gewährleistung",
        "Einbehalte, Bürgschaften, Aval — Restlaufzeit steuert die Liquiditätsbindung nach Abnahme",
        10,
        {1: 10, 2: 14, 3: 12, 4: 12, 5: 12, 6: 14, 7: 12, 8: 12, 9: 32, 10: 16},
    )
    write_headers(ws, GW_H)
    kinds = ["text", "text", "text", "date", "date", "eur", "text", "text", "text", "calc:qty"]
    for i, src in enumerate(sd.GEWAEHR):
        r = 4 + i
        vals = list(src) + [f'=IF(A{r}="","",MAX(0,(E{r}-TODAY())/365))']
        write_row(ws, r, vals, kinds)
    for r in range(4 + len(sd.GEWAEHR), MAX_ROWS + 1):
        write_row(ws, r, [""] * 9 + [f'=IF(A{r}="","",MAX(0,(E{r}-TODAY())/365))'], kinds)
    add_table(ws, "tblGewaehr", f"A3:J{MAX_ROWS}")
    dropdown(ws, f"B4:B{MAX_ROWS}", "ProjektNr")
    dropdown(ws, f"G4:G{MAX_ROWS}", '"Einbehalt,Bürgschaft,Aval,Gewährleistung"')
    ws.freeze_panes = "A4"


def add_all_finanzen(wb: Workbook) -> None:
    add_ar(wb)
    add_er(wb)
    add_zahlungen(wb)
    add_opos(wb)
    add_mahnwesen(wb)
    add_gewaehr(wb)
