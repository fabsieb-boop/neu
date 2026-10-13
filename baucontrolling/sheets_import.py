"""Importblatt und Bedienungsanleitung."""

from openpyxl.styles import Alignment, Font
from openpyxl.workbook.workbook import Workbook

from baucontrolling.layout import new_sheet
from baucontrolling.styles import (
    ALIGN_L,
    FILL_NAVY,
    FILL_WHITE,
    FONT_BODY,
    FONT_H2,
    FONT_H3,
    FONT_HEAD,
    FONT_SMALL,
    FONT_WHITE,
    THIN,
    style_input,
)


def add_import(wb: Workbook) -> None:
    ws = new_sheet(
        wb,
        "30_Import",
        "Automatische Importe",
        "CSV-Vorlagen im Ordner templates/ · Python-Skript scripts/import_csv.py · optional Power Query",
        8,
        {1: 26, 2: 22, 3: 18, 4: 18, 5: 18, 6: 28, 7: 18, 8: 18},
    )
    ws["A4"] = "1. Empfohlener Weg: Python-Import (ohne Excel-Makros)"
    ws["A4"].font = FONT_H2
    ws.merge_cells("A5:F8")
    ws["A5"] = (
        "Im Projektordner:\n"
        "python scripts/import_csv.py --workbook Baucontrolling_Kaufmann.xlsx --typ bank --datei templates/bankumsatz.csv\n"
        "Typen: kunden | lieferanten | personal | projekte | stunden | material | ar | er | zahlungen | bank | datev\n"
        "Das Skript hängt Zeilen an die richtige Tabelle an und vergibt laufende Nummern, wenn die Nr. leer ist."
    )
    ws["A5"].alignment = Alignment(wrap_text=True, vertical="top")

    ws["A10"] = "2. Power Query (Excel: Daten → Aus Text/CSV)"
    ws["A10"].font = FONT_H2
    ws.merge_cells("A11:F14")
    ws["A11"] = (
        "Datei wählen → Trennzeichen Semikolon → Laden in Tabelle. "
        "Anschließend Abfrage so benennen wie unten. Beim nächsten Mal: Daten → Alle aktualisieren. "
        "Spaltenüberschriften müssen exakt zur Ziel-Tabelle passen (siehe Vorlagen)."
    )
    ws["A11"].alignment = Alignment(wrap_text=True, vertical="top")

    ws["A16"] = "Zielblatt"
    ws["B16"] = "CSV-Vorlage"
    ws["C16"] = "Typ-Flag"
    ws["D16"] = "Trennzeichen"
    ws["E16"] = "Encoding"
    ws["F16"] = "Bemerkung"
    for col in range(1, 7):
        ws.cell(16, col).fill = FILL_NAVY
        ws.cell(16, col).font = FONT_HEAD
        ws.cell(16, col).border = THIN
    mapping = [
        ("02_Kunden", "templates/kunden.csv", "kunden", ";", "utf-8-sig", "Debitorenstamm"),
        ("03_Lieferanten", "templates/lieferanten.csv", "lieferanten", ";", "utf-8-sig", "Kreditorenstamm"),
        ("04_Personal", "templates/personal.csv", "personal", ";", "utf-8-sig", "Sätze überschreiben bestehendes Nr."),
        ("05_Projekte", "templates/projekte.csv", "projekte", ";", "utf-8-sig", "nur Stammdaten-Spalten A–T"),
        ("10_Stunden", "templates/stunden.csv", "stunden", ";", "utf-8-sig", "aus Terminal / App exportieren"),
        ("11_Material", "templates/material.csv", "material", ";", "utf-8-sig", "Lieferscheinpositionen"),
        ("13_Ausgangsrechnungen", "templates/ausgangsrechnungen.csv", "ar", ";", "utf-8-sig", "keine Formelspalten importieren"),
        ("14_Eingangsrechnungen", "templates/eingangsrechnungen.csv", "er", ";", "utf-8-sig", "DATEV können gemappt werden"),
        ("15_Zahlungen", "templates/zahlungen.csv", "zahlungen", ";", "utf-8-sig", "manuelle Zahlungen"),
        ("15_Zahlungen", "templates/bankumsatz.csv", "bank", ";", "utf-8-sig", "Sparkasse/VR-CAMT-ähnlich"),
        ("15_Zahlungen", "templates/datev_buchungen.csv", "datev", ";", "utf-8-sig", "Buchungsstapel EXTFähig"),
    ]
    for i, row in enumerate(mapping):
        for c, val in enumerate(row, start=1):
            ws.cell(17 + i, c, val)
            style_input(ws.cell(17 + i, c))

    ws["A30"] = "3. Bank-CSV Spaltenzuordnung (Typ bank)"
    ws["A30"].font = FONT_H2
    bank_map = [
        ("Datum", "Buchungstag oder Valuta"),
        ("Betrag", "positiv = Eingang, negativ = Ausgang — Skript dreht das Vorzeichen in Art um"),
        ("Gegenkonto / Partner", "wird Partner"),
        ("Verwendungszweck", "Skript sucht AR-xxxx oder ER-xxxx und setzt Beleg"),
        ("BIC/Bank", "Weg = Bank"),
    ]
    ws["A31"] = "CSV-Spalte"
    ws["B31"] = "Bedeutung"
    ws["A31"].fill = FILL_NAVY
    ws["B31"].fill = FILL_NAVY
    ws["A31"].font = FONT_HEAD
    ws["B31"].font = FONT_HEAD
    for i, (a, b) in enumerate(bank_map):
        ws.cell(32 + i, 1, a)
        ws.cell(32 + i, 2, b)
        style_input(ws.cell(32 + i, 1))
        style_input(ws.cell(32 + i, 2))

    ws["A39"] = "4. DATEV-Buchungsstapel (Typ datev)"
    ws.merge_cells("A40:F42")
    ws["A40"] = (
        "Erwartete Spalten: Umsatz (ohne Soll/Haben-KZ), Soll/Haben-KZ, Konto, Gegenkonto, Belegdatum, "
        "Buchungstext, Belegfeld1. Konten 8400–8599 bzw. 4400–4499 werden als Ausgangsrechnung grob erkannt, "
        "Eingangsrechnungen über Aufwandskonten 3000–7999. Für den produktiven Betrieb Kontenrahmen in scripts/import_csv.py anpassen."
    )
    ws["A40"].alignment = Alignment(wrap_text=True, vertical="top")

    ws["A44"] = "Letzter Import (manuell notieren)"
    ws["A45"] = "Datum"
    ws["B45"] = "Datei"
    ws["C45"] = "Typ"
    ws["D45"] = "Zeilen"
    ws["E45"] = "wer"
    for col in range(1, 6):
        ws.cell(45, col).fill = FILL_NAVY
        ws.cell(45, col).font = FONT_HEAD
    for r in range(46, 56):
        for c in range(1, 6):
            style_input(ws.cell(r, c), "date" if c == 1 else "text")


def add_anleitung(wb: Workbook) -> None:
    ws = new_sheet(
        wb,
        "31_Anleitung",
        "Bedienung für Inhaber, Bauleitung und Buchhaltung",
        "Diese Mappe ersetzt kein DATEV-Rechnungswesen, steuert aber Liquidität, Projekte, Forecast und Formulare im Alltag",
        6,
        {1: 28, 2: 88, 3: 16, 4: 16, 5: 16, 6: 16},
    )
    blocks = [
        ("Tägliche Arbeit Bauleitung", "Stunden auf 10_Stunden buchen, Material auf 11_Material, Bautagesbericht über 28_Form_Bautage drucken (Projekt + Datum wählen). Nachträge sofort auf 07_Nachtraege mit Status angeboten / in Prüfung / beauftragt."),
        ("Wöchentlich Buchhaltung", "Eingangsrechnungen auf 14_ erfassen, Zahlungen importieren (Bank-CSV), OPOS 16_ prüfen, Mahnwesen 17_ eine Stufe weiter, Dashboard 24_ mit der GF besprechen. Gelbe Zellen in 20_Liquiditaet auf bekannte Abschläge und Lohnläufe aktualisieren."),
        ("Monatlich Kaufmann / GF", "Forecast 21_ rollierend machen (Ist bleibt Formel, Plan in den gelben Planzeilen). BWA 22_ und Soll-Ist 23_ gegen Nachkalkulation lesen. Einbehalte auf 18_ nachziehen. Angebote ohne Entscheidung in 08_ bereinigen."),
        ("Projekt neu", "Kunde anlegen → Angebot 08_ → bei Zuschlag Auftrag 09_ und Projekt 05_ → LV-Positionen 06_ → Budgets in 05_ (Material, NU, Lohn, Gerät, GK). Ab dann laufen DB, Fertigstellung und OP automatisch."),
        ("Rechnung schreiben", "Zeile in 13_Ausgangsrechnungen anlegen (Typ Abschlag/Schluss, Netto, Einbehalt %). Blatt 26_Form_Rechnung öffnen, Nummer wählen, drucken. Zahlung später per Import — Offen und Status folgen der Bank."),
        ("Was die Mappe bewusst nicht ist", "Kein Ersatz für Lohnabrechnung, Anlagenbuchhaltung oder die USt-Voranmeldung. Keine GoBD-revisionssichere Buchführung. Belege bleiben im DATEV/Unternehmen-online; diese Datei ist das Steuerungs-Cockpit."),
        ("Ampeln", "OK = unkritisch. ENG = unter 125 % der Mindestliquidität. ENGPASS = unter Mindestliquidität. PRÜFEN im Soll-Ist = Marge < 8 %, Kosten über Budget oder OP > 50 T€."),
        ("Datensicherheit", "Datei auf dem Firmenlaufwerk versionieren, monatlich unter Archiv/JJJJ-MM speichern. Personal- und Bankdaten nicht in ungesicherte Cloud-Ordner legen."),
    ]
    r = 4
    for title, text in blocks:
        ws.cell(r, 1, title).font = FONT_H3
        ws.merge_cells(start_row=r, start_column=2, end_row=r + 2, end_column=6)
        cell = ws.cell(r, 2, text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.font = FONT_BODY
        cell.fill = FILL_WHITE
        cell.border = THIN
        r += 4
    ws.freeze_panes = "A4"


def add_all_docs(wb: Workbook) -> None:
    add_import(wb)
    add_anleitung(wb)
