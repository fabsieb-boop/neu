"""Farben, Schriften und Zellstile für das Baucontrolling-Workbook."""

from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Side
from openpyxl.styles.numbers import FORMAT_NUMBER_COMMA_SEPARATED1
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

NAVY = "1B365D"
STEEL = "2C5282"
INK = "1A1A1A"
MUTED = "5C5C5C"
PAPER = "F7F4EE"
CREAM = "EFEBE3"
WHITE = "FFFFFF"
LINE = "D4CFC4"
GOLD = "B8860B"
GREEN = "2F6B4F"
GREEN_SOFT = "E4F0E8"
RED = "8B2E2E"
RED_SOFT = "F6E4E4"
AMBER = "8A6A12"
AMBER_SOFT = "F6EED8"
BLUE_SOFT = "E4EAF2"

THIN = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)
THIN_NAVY = Border(
    left=Side(style="thin", color=NAVY),
    right=Side(style="thin", color=NAVY),
    top=Side(style="thin", color=NAVY),
    bottom=Side(style="thin", color=NAVY),
)

FONT_TITLE = Font(name="Calibri", size=20, bold=True, color=NAVY)
FONT_H2 = Font(name="Calibri", size=14, bold=True, color=NAVY)
FONT_H3 = Font(name="Calibri", size=11, bold=True, color=STEEL)
FONT_HEAD = Font(name="Calibri", size=9, bold=True, color=WHITE)
FONT_BODY = Font(name="Calibri", size=10, color=INK)
FONT_SMALL = Font(name="Calibri", size=8, color=MUTED)
FONT_KPI = Font(name="Calibri", size=18, bold=True, color=NAVY)
FONT_WHITE = Font(name="Calibri", size=10, bold=True, color=WHITE)

FILL_NAVY = PatternFill("solid", fgColor=NAVY)
FILL_STEEL = PatternFill("solid", fgColor=STEEL)
FILL_PAPER = PatternFill("solid", fgColor=PAPER)
FILL_CREAM = PatternFill("solid", fgColor=CREAM)
FILL_WHITE = PatternFill("solid", fgColor=WHITE)
FILL_GOLD = PatternFill("solid", fgColor=GOLD)
FILL_GREEN = PatternFill("solid", fgColor=GREEN)
FILL_RED = PatternFill("solid", fgColor=RED)
FILL_GREEN_SOFT = PatternFill("solid", fgColor=GREEN_SOFT)
FILL_RED_SOFT = PatternFill("solid", fgColor=RED_SOFT)
FILL_AMBER_SOFT = PatternFill("solid", fgColor=AMBER_SOFT)
FILL_BLUE_SOFT = PatternFill("solid", fgColor=BLUE_SOFT)

EUR = '#,##0.00 "€"'
EUR0 = '#,##0 "€"'
PCT = "0.0%"
DATE = "DD.MM.YYYY"
QTY = "#,##0.00"
HRS = '#,##0.00 "h"'

ALIGN_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_L = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_R = Alignment(horizontal="right", vertical="center")


def apply_workbook_defaults(wb: Workbook) -> None:
    if "Normal" not in wb.named_styles:
        normal = NamedStyle(name="Normal")
        normal.font = FONT_BODY
        wb.add_named_style(normal)


def paint_sheet(ws: Worksheet, color: str = PAPER, rows: int = 80, cols: int = 20) -> None:
    fill = PatternFill("solid", fgColor=color)
    for row in ws.iter_rows(min_row=1, max_row=rows, min_col=1, max_col=cols):
        for cell in row:
            cell.fill = fill
            cell.font = FONT_BODY


def set_widths(ws: Worksheet, widths: dict[int, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def header_bar(ws: Worksheet, title: str, subtitle: str, last_col: int = 12) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    c1 = ws.cell(1, 1, title)
    c1.font = Font(name="Calibri", size=18, bold=True, color=WHITE)
    c1.fill = FILL_NAVY
    c1.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    c2 = ws.cell(2, 1, subtitle)
    c2.font = Font(name="Calibri", size=10, color=WHITE)
    c2.fill = FILL_STEEL
    c2.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 18
    for col in range(1, last_col + 1):
        ws.cell(1, col).fill = FILL_NAVY
        ws.cell(2, col).fill = FILL_STEEL


def section_label(ws: Worksheet, row: int, col: int, text: str, span: int = 4) -> None:
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)
    cell = ws.cell(row, col, text)
    cell.font = FONT_WHITE
    cell.fill = FILL_STEEL
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for i in range(span):
        ws.cell(row, col + i).fill = FILL_STEEL


def style_header_row(ws: Worksheet, row: int, cols: int, start: int = 1) -> None:
    ws.row_dimensions[row].height = 28
    for col in range(start, start + cols):
        cell = ws.cell(row, col)
        cell.font = FONT_HEAD
        cell.fill = FILL_NAVY
        cell.alignment = ALIGN_C
        cell.border = THIN


def style_input(cell, kind: str = "text") -> None:
    cell.font = FONT_BODY
    cell.fill = FILL_WHITE
    cell.border = THIN
    cell.alignment = ALIGN_L
    if kind == "eur":
        cell.number_format = EUR
        cell.alignment = ALIGN_R
    elif kind == "pct":
        cell.number_format = PCT
        cell.alignment = ALIGN_R
    elif kind == "date":
        cell.number_format = DATE
        cell.alignment = ALIGN_C
    elif kind == "qty":
        cell.number_format = QTY
        cell.alignment = ALIGN_R
    elif kind == "hrs":
        cell.number_format = HRS
        cell.alignment = ALIGN_R
    elif kind == "int":
        cell.number_format = "0"
        cell.alignment = ALIGN_C


def style_calc(cell, kind: str = "eur") -> None:
    cell.font = Font(name="Calibri", size=10, color=STEEL)
    cell.fill = FILL_BLUE_SOFT
    cell.border = THIN
    cell.alignment = ALIGN_R
    if kind == "eur":
        cell.number_format = EUR
    elif kind == "pct":
        cell.number_format = PCT
    elif kind == "date":
        cell.number_format = DATE
        cell.alignment = ALIGN_C
    elif kind == "qty":
        cell.number_format = QTY
    elif kind == "hrs":
        cell.number_format = HRS
    elif kind == "text":
        cell.alignment = ALIGN_L


def zebra(ws: Worksheet, start: int, end: int, cols: int, col0: int = 1) -> None:
    for r in range(start, end + 1):
        fill = FILL_CREAM if (r - start) % 2 else FILL_WHITE
        for c in range(col0, col0 + cols):
            cell = ws.cell(r, c)
            if cell.fill.fgColor is None or cell.fill.fgColor.rgb in {None, "00000000", WHITE, PAPER}:
                cell.fill = fill
            cell.border = THIN
            cell.font = FONT_BODY


def kpi_box(ws: Worksheet, row: int, col: int, label: str, formula: str, kind: str = "eur", span: int = 2) -> None:
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)
    ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 2, end_column=col + span - 1)
    lab = ws.cell(row, col, label)
    lab.font = Font(name="Calibri", size=8, bold=True, color=WHITE)
    lab.fill = FILL_STEEL
    lab.alignment = ALIGN_C
    val = ws.cell(row + 1, col, formula)
    val.font = FONT_KPI
    val.fill = FILL_WHITE
    val.border = THIN_NAVY
    val.alignment = Alignment(horizontal="center", vertical="center")
    if kind == "eur":
        val.number_format = EUR0
    elif kind == "pct":
        val.number_format = PCT
    elif kind == "int":
        val.number_format = "0"
    for c in range(col, col + span):
        ws.cell(row, c).fill = FILL_STEEL
        ws.cell(row + 1, c).border = THIN_NAVY
        ws.cell(row + 2, c).border = THIN_NAVY
        ws.cell(row + 2, c).fill = FILL_WHITE
    ws.row_dimensions[row + 1].height = 22
    ws.row_dimensions[row + 2].height = 10


def freeze(ws: Worksheet, cell: str = "A4") -> None:
    ws.freeze_panes = cell
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.horizontalCentered = True
    ws.print_title_rows = "1:3"
    ws.oddHeader.left.text = "&A"
    ws.oddHeader.right.text = "&D"
    ws.oddFooter.left.text = "Bergmann Hoch- und Tiefbau GmbH · vertraulich"
    ws.oddFooter.right.text = "Seite &P / &N"
