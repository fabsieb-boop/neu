"""Projekte, LV, Nachträge, Angebote, Aufträge."""

from openpyxl.workbook.workbook import Workbook

from baucontrolling import sample_data as sd
from baucontrolling.layout import (
    MAX_ROWS,
    add_table,
    dropdown,
    fill_blank_rows,
    named,
    new_sheet,
    status_colors,
    write_headers,
    write_row,
)

PROJ_H = [
    "Nr", "Name", "Kunde", "Art", "Vertrag", "Status", "Start", "Ende Plan", "Ende Ist",
    "Auftrag netto", "Budget Material", "Budget NU", "Budget Lohn", "Budget Gerät", "Budget GK",
    "Einbehalt %", "Bauleiter", "Polier", "Ort", "Notiz",
    "Nachträge beauftr.", "Auftrag aktuell", "Ist Leistung", "Ist Material", "Ist NU",
    "Ist Lohn", "Ist Kosten", "DB Ist", "Marge %", "Fertigstellung", "Restumsatz", "OP Debitor",
]
PROJ_IN = [
    "text", "text", "text", "text", "text", "text", "date", "date", "date",
    "eur", "eur", "eur", "eur", "eur", "eur",
    "pct", "text", "text", "text", "text",
]
# calc from col 21
PROJ_CALC = ["calc:eur", "calc:eur", "calc:eur", "calc:eur", "calc:eur", "calc:eur", "calc:eur", "calc:eur", "calc:pct", "calc:pct", "calc:eur", "calc:eur"]

LV_H = ["Pos-Nr", "Projekt", "OZ", "Bezeichnung", "ME", "Menge", "EP netto", "Faktor", "Gewerk", "Status", "Gesamt netto"]
NACH_H = ["Nr", "Projekt", "Datum", "Bezeichnung", "Betrag netto", "Status", "Kunde", "Begründung", "Auftragswirksam"]
ANG_H = ["Nr", "Datum", "Kunde", "Projekt", "Bezeichnung", "Netto", "AGK %", "W+G %", "Status", "Entscheidung", "Bindefrist Tage", "Angebotssumme inkl. Zuschlag", "Hit-Rate-Hilfe"]
AUF_H = ["Nr", "Angebot", "Projekt", "Kunde", "Datum", "Auftrag netto", "Vertragsart", "Einbehalt %", "Status", "Notiz", "Nachträge", "Auftrag aktuell"]


def _proj_formulas(r: int) -> list:
    return [
        f'=IF(A{r}="","",SUMIFS(\'07_Nachtraege\'!$E:$E,\'07_Nachtraege\'!$B:$B,A{r},\'07_Nachtraege\'!$F:$F,"beauftragt"))',
        f'=IF(A{r}="","",J{r}+U{r})',
        f'=IF(A{r}="","",SUMIFS(\'13_Ausgangsrechnungen\'!$G:$G,\'13_Ausgangsrechnungen\'!$C:$C,A{r}))',
        f'=IF(A{r}="","",SUMIFS(\'11_Material\'!$L:$L,\'11_Material\'!$C:$C,A{r}))',
        f'=IF(A{r}="","",SUMIFS(\'14_Eingangsrechnungen\'!$F:$F,\'14_Eingangsrechnungen\'!$C:$C,A{r},\'14_Eingangsrechnungen\'!$E:$E,"Nachunternehmer"))',
        f'=IF(A{r}="","",SUMIFS(\'10_Stunden\'!$K:$K,\'10_Stunden\'!$C:$C,A{r}))',
        f'=IF(A{r}="","",X{r}+Y{r}+Z{r}+SUMIFS(\'14_Eingangsrechnungen\'!$F:$F,\'14_Eingangsrechnungen\'!$C:$C,A{r},\'14_Eingangsrechnungen\'!$E:$E,"Gerät")+SUMIFS(\'14_Eingangsrechnungen\'!$F:$F,\'14_Eingangsrechnungen\'!$C:$C,A{r},\'14_Eingangsrechnungen\'!$E:$E,"Dienstleistung"))',
        f'=IF(A{r}="","",W{r}-AA{r})',
        f'=IF(OR(A{r}="",W{r}=0),"",AB{r}/W{r})',
        f'=IF(OR(A{r}="",V{r}=0),"",W{r}/V{r})',
        f'=IF(A{r}="","",MAX(0,V{r}-W{r}))',
        f'=IF(A{r}="","",SUMIFS(\'13_Ausgangsrechnungen\'!$R:$R,\'13_Ausgangsrechnungen\'!$C:$C,A{r}))',
    ]


def add_projekte(wb: Workbook) -> None:
    widths = {i: 14 for i in range(1, 33)}
    widths.update({1: 14, 2: 26, 3: 12, 20: 22, 19: 14})
    ws = new_sheet(
        wb,
        "05_Projekte",
        "Projektstamm und Deckungsbeitrag",
        "Auftrag, Nachträge, Ist-Leistung und Kosten werden automatisch verdichtet — Ampel über Marge und OP",
        32,
        widths,
    )
    write_headers(ws, PROJ_H)
    kinds = PROJ_IN + PROJ_CALC
    for i, src in enumerate(sd.PROJEKTE):
        r = 4 + i
        values = list(src) + _proj_formulas(r)
        write_row(ws, r, values, kinds)
    for r in range(4 + len(sd.PROJEKTE), MAX_ROWS + 1):
        empty = [""] * 20 + _proj_formulas(r)
        write_row(ws, r, empty, kinds)
    add_table(ws, "tblProjekte", f"A3:AF{MAX_ROWS}")
    # wait, last col is 32 = AF? A=1 ... 32 = AF yes
    # I used AF in table but headers have 32 cols. Let me count: 32 headers, last is AF.
    dropdown(ws, f"C4:C{MAX_ROWS}", "KundenNr")
    dropdown(ws, f"D4:D{MAX_ROWS}", "lstProjektArt")
    dropdown(ws, f"E4:E{MAX_ROWS}", "lstVertrag")
    dropdown(ws, f"F4:F{MAX_ROWS}", "lstProjektStatus")
    dropdown(ws, f"Q4:Q{MAX_ROWS}", "PersNr")
    dropdown(ws, f"R4:R{MAX_ROWS}", "PersNr")
    status_colors(ws, 6, 4, MAX_ROWS)
    named(wb, "ProjektNr", f"'05_Projekte'!$A$4:$A${MAX_ROWS}")
    named(wb, "ProjektName", f"'05_Projekte'!$B$4:$B${MAX_ROWS}")
    ws.freeze_panes = "A4"


def add_lv(wb: Workbook) -> None:
    ws = new_sheet(wb, "06_LV", "Leistungsverzeichnis", "Einheitspreisvertrag: Menge × EP × Faktor = Positionssumme, verdichtet je Projekt", 11, {1: 12, 2: 14, 3: 10, 4: 32, 5: 12, 6: 12, 7: 14, 8: 10, 9: 14, 10: 12, 11: 16})
    write_headers(ws, LV_H)
    kinds = ["text", "text", "text", "text", "text", "qty", "eur", "qty", "text", "text", "calc:eur"]
    for i, src in enumerate(sd.LV):
        r = 4 + i
        vals = list(src) + [f'=IF(A{r}="","",F{r}*G{r}*H{r})']
        write_row(ws, r, vals, kinds)
    for r in range(4 + len(sd.LV), MAX_ROWS + 1):
        write_row(ws, r, [""] * 10 + [f'=IF(A{r}="","",F{r}*G{r}*H{r})'], kinds)
    add_table(ws, "tblLV", f"A3:K{MAX_ROWS}")
    dropdown(ws, f"B4:B{MAX_ROWS}", "ProjektNr")
    dropdown(ws, f"J4:J{MAX_ROWS}", "lstProjektStatus")
    ws.freeze_panes = "A4"


def add_nachtraege(wb: Workbook) -> None:
    ws = new_sheet(wb, "07_Nachtraege", "Nachträge und Claims", "Nur Status „beauftragt“ erhöht den aktuellen Auftragswert im Projektstamm", 9, {1: 12, 2: 14, 3: 12, 4: 34, 5: 14, 6: 14, 7: 12, 8: 32, 9: 16})
    write_headers(ws, NACH_H)
    kinds = ["text", "text", "date", "text", "eur", "text", "text", "text", "calc:eur"]
    for i, src in enumerate(sd.NACHTRAEGE):
        r = 4 + i
        vals = list(src) + [f'=IF(A{r}="","",IF(F{r}="beauftragt",E{r},0))']
        write_row(ws, r, vals, kinds)
    for r in range(4 + len(sd.NACHTRAEGE), MAX_ROWS + 1):
        write_row(ws, r, [""] * 8 + [f'=IF(A{r}="","",IF(F{r}="beauftragt",E{r},0))'], kinds)
    add_table(ws, "tblNachtraege", f"A3:I{MAX_ROWS}")
    dropdown(ws, f"B4:B{MAX_ROWS}", "ProjektNr")
    dropdown(ws, f"F4:F{MAX_ROWS}", '"angeboten,in Prüfung,beauftragt,abgelehnt,zurückgezogen"')
    ws.freeze_panes = "A4"


def add_angebote(wb: Workbook) -> None:
    ws = new_sheet(wb, "08_Angebote", "Angebotsverfolgung", "Hit-Rate, Bindefrist und Zuschläge AGK / Wagnis+Gewinn aus Stammdaten vorbelegt", 13, {1: 12, 2: 12, 3: 12, 4: 14, 5: 28, 6: 14, 7: 10, 8: 10, 9: 14, 10: 14, 11: 12, 12: 20, 13: 16})
    write_headers(ws, ANG_H)
    kinds = ["text", "date", "text", "text", "text", "eur", "pct", "pct", "text", "date", "int", "calc:eur", "calc:text"]
    for i, src in enumerate(sd.ANGEBOTE):
        r = 4 + i
        vals = list(src) + [
            f'=IF(A{r}="","",F{r}*(1+G{r})*(1+H{r}))',
            f'=IF(A{r}="","",IF(I{r}="angenommen","Auftrag",IF(AND(I{r}="offen",TODAY()>B{r}+K{r}),"Bindefrist!","")))',
        ]
        write_row(ws, r, vals, kinds)
    for r in range(4 + len(sd.ANGEBOTE), MAX_ROWS + 1):
        write_row(
            ws,
            r,
            [""] * 11
            + [
                f'=IF(A{r}="","",F{r}*(1+G{r})*(1+H{r}))',
                f'=IF(A{r}="","",IF(I{r}="angenommen","Auftrag",IF(AND(I{r}="offen",TODAY()>B{r}+K{r}),"Bindefrist!","")))',
            ],
            kinds,
        )
    add_table(ws, "tblAngebote", f"A3:M{MAX_ROWS}")
    dropdown(ws, f"C4:C{MAX_ROWS}", "KundenNr")
    dropdown(ws, f"D4:D{MAX_ROWS}", "ProjektNr")
    dropdown(ws, f"I4:I{MAX_ROWS}", "lstAngebotStatus")
    named(wb, "AngebotNr", f"'08_Angebote'!$A$4:$A${MAX_ROWS}")
    ws.freeze_panes = "A4"


def add_auftraege(wb: Workbook) -> None:
    ws = new_sheet(wb, "09_Auftraege", "Auftragsbestand", "Vertragsdaten und aktueller Auftragswert inkl. beauftragter Nachträge", 12, {1: 12, 2: 12, 3: 14, 4: 12, 5: 12, 6: 16, 7: 12, 8: 12, 9: 14, 10: 28, 11: 14, 12: 16})
    write_headers(ws, AUF_H)
    kinds = ["text", "text", "text", "text", "date", "eur", "text", "pct", "text", "text", "calc:eur", "calc:eur"]
    for i, src in enumerate(sd.AUFTRAEGE):
        r = 4 + i
        vals = list(src) + [
            f'=IF(A{r}="","",SUMIFS(\'07_Nachtraege\'!$E:$E,\'07_Nachtraege\'!$B:$B,C{r},\'07_Nachtraege\'!$F:$F,"beauftragt"))',
            f'=IF(A{r}="","",F{r}+K{r})',
        ]
        write_row(ws, r, vals, kinds)
    for r in range(4 + len(sd.AUFTRAEGE), MAX_ROWS + 1):
        write_row(
            ws,
            r,
            [""] * 10
            + [
                f'=IF(A{r}="","",SUMIFS(\'07_Nachtraege\'!$E:$E,\'07_Nachtraege\'!$B:$B,C{r},\'07_Nachtraege\'!$F:$F,"beauftragt"))',
                f'=IF(A{r}="","",F{r}+K{r})',
            ],
            kinds,
        )
    add_table(ws, "tblAuftraege", f"A3:L{MAX_ROWS}")
    dropdown(ws, f"B4:B{MAX_ROWS}", "AngebotNr")
    dropdown(ws, f"C4:C{MAX_ROWS}", "ProjektNr")
    dropdown(ws, f"D4:D{MAX_ROWS}", "KundenNr")
    dropdown(ws, f"G4:G{MAX_ROWS}", "lstVertrag")
    ws.freeze_panes = "A4"


def add_all_projekte(wb: Workbook) -> None:
    add_projekte(wb)
    add_lv(wb)
    add_nachtraege(wb)
    add_angebote(wb)
    add_auftraege(wb)
