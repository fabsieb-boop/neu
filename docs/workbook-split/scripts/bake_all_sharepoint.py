#!/usr/bin/env python3
"""Bake Excel formula caches into all SharePoint workbooks.

Excel Online / SharePoint File Viewer does not reliably resolve cross-workbook
links → #WERT! / blank KPIs. This script replaces formula cells with the last
calculated values from TVF_Controlling_Cockpit_Import.xlsx (literals).
Hyperlink formulas are kept.
"""

from __future__ import annotations

import shutil
import zipfile
import warnings
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

warnings.filterwarnings("ignore")

ROOT = Path(__file__).resolve().parents[1]
STORE_DOCS = Path("/cursor/stores/bc-d9b5c533-1abd-48a4-a0b7-3919e592c358/docs")
DATEIEN = ROOT / "dateien"

IMPORT_CANDIDATES = [
    STORE_DOCS / "TVF_Controlling_Cockpit_Import.xlsx",
    ROOT.parent / "TVF_Controlling_Cockpit_Import.xlsx",
]

HUB_SHEETS = (
    "Cockpit",
    "Steuerung",
    "GuV",
    "Forecast",
    "Monatsroll",
    "FC Eingabe",
    "Import Monatsdaten",
    "D Finanzen",
)

KEEP_MARKERS = ("HYPERLINK",)


def find_import() -> Path:
    for p in IMPORT_CANDIDATES:
        if p.exists():
            return p
    raise FileNotFoundError("TVF_Controlling_Cockpit_Import.xlsx fehlt")


def is_keep_formula(value: object) -> bool:
    if not isinstance(value, str) or not value.startswith("="):
        return False
    upper = value.upper()
    return any(m in upper for m in KEEP_MARKERS)


def bake_sheet(ws_target, ws_val, ws_src) -> tuple[int, int, int]:
    baked = kept = emptied = 0
    max_row = max(ws_target.max_row or 1, ws_src.max_row or 1)
    max_col = max(ws_target.max_column or 1, ws_src.max_column or 1)
    for row in ws_target.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for cell in row:
            src = ws_src[cell.coordinate].value
            cur = cell.value
            formula = None
            if isinstance(cur, str) and cur.startswith("="):
                formula = cur
            elif isinstance(src, str) and src.startswith("="):
                formula = src
            if formula is None:
                continue
            if is_keep_formula(formula):
                kept += 1
                continue
            cached = ws_val[cell.coordinate].value
            if cached is None:
                cell.value = None
                emptied += 1
            else:
                cell.value = cached
                baked += 1
    return baked, kept, emptied


def bake_workbook(path: Path, wb_val, wb_src, sheet_filter: tuple[str, ...] | None = None) -> dict:
    wb = load_workbook(path, data_only=False)
    stats = {}
    for name in list(wb.sheetnames):
        if sheet_filter and name not in sheet_filter:
            continue
        if name not in wb_val.sheetnames:
            stats[name] = {"baked": 0, "kept": 0, "emptied": 0, "skip": "not in import"}
            continue
        b, k, e = bake_sheet(wb[name], wb_val[name], wb_src[name])
        stats[name] = {"baked": b, "kept": k, "emptied": e}
    # Align Steuerung year/month with Import when present
    if "Steuerung" in wb.sheetnames and "Steuerung" in wb_val.sheetnames:
        for coord in ("F6", "F7"):
            v = wb_val["Steuerung"][coord].value
            if v is not None:
                wb["Steuerung"][coord] = v
    if "Cockpit" in wb.sheetnames:
        ws = wb["Cockpit"]
        if "Steuerung" in wb.sheetnames:
            year = wb["Steuerung"]["F6"].value
            month = wb["Steuerung"]["F7"].value
            if isinstance(year, (int, float)) and isinstance(month, (int, float)):
                ws["C9"] = f"Berichtszeitraum: 01–{int(month):02d} / {int(year)}"
        ws["C65"] = (
            "Kennzahlen sind für SharePoint fest geschrieben (kein Excel-Online-Neurechnen). "
            "Externe Verknüpfungen wurden durch Werte ersetzt."
        )
        ws["C66"] = (
            "Neu berechnen: in Excel Desktop öffnen, rechnen, speichern – "
            "oder scripts/bake_all_sharepoint.py erneut ausführen und Dateien neu hochladen."
        )
        for coord in ("C65", "C66"):
            ws[coord].font = Font(name="Calibri", size=9, italic=True, color="666666")
    wb.save(path)
    return stats


def rebuild_hub(wb_val, wb_src) -> Path:
    hub = ROOT / "Cockpit-Hub-SharePoint.xlsx"
    stats = bake_workbook(hub, wb_val, wb_src, sheet_filter=HUB_SHEETS)
    total = sum(s.get("baked", 0) for s in stats.values())
    print(f"Hub: {total} Zellen gebacken")
    for dest in (
        ROOT / "Cockpit.xlsx",
        DATEIEN / "Cockpit.xlsx",
        ROOT.parent / "Cockpit.xlsx",
        ROOT.parent / "Cockpit-Hub-SharePoint.xlsx",
        STORE_DOCS / "Cockpit.xlsx",
        STORE_DOCS / "Cockpit-Hub-SharePoint.xlsx",
        STORE_DOCS / "workbook-split" / "Cockpit.xlsx",
        STORE_DOCS / "workbook-split" / "Cockpit-Hub-SharePoint.xlsx",
    ):
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(hub, dest)
    return hub


def bake_dateien(wb_val, wb_src) -> None:
    files = sorted(DATEIEN.glob("*.xlsx"))
    for path in files:
        if path.name == "Cockpit.xlsx":
            continue  # hub copy handled separately
        stats = bake_workbook(path, wb_val, wb_src)
        baked = sum(s.get("baked", 0) for s in stats.values())
        emptied = sum(s.get("emptied", 0) for s in stats.values())
        print(f"  {path.name}: baked={baked} emptied={emptied}")
        # mirror into store
        store_dest = STORE_DOCS / "workbook-split" / "dateien" / path.name
        store_dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(path, store_dest)


def write_zip(zip_path: Path, files: list[tuple[Path, str]]) -> None:
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for src, arc in files:
            if src.exists():
                zf.write(src, arcname=arc)


def refresh_packages() -> None:
    cockpit = ROOT / "Cockpit.xlsx"
    hub = ROOT / "Cockpit-Hub-SharePoint.xlsx"

    dl = ROOT / "TVF_Cockpit_Download.zip"
    write_zip(
        dl,
        [
            (cockpit, "Cockpit.xlsx"),
        ],
    )
    with zipfile.ZipFile(dl, "a") as zf:
        zf.writestr(
            "README.txt",
            "Cockpit.xlsx = Hub mit eingebackenen Kennzahlen.\n"
            "Nach TVF-Controlling-Cockpit/Cockpit.xlsx hochladen (ersetzen).\n"
            "Seite veröffentlichen + Hard-Reload (Cmd+Shift+R).\n"
            "Einzeldateien (Stunden, KLR, …) liegen in dateien/ bzw. TVF_Cockpit_Dateien.zip.\n",
        )

    dateien_zip = ROOT / "TVF_Cockpit_Dateien.zip"
    with zipfile.ZipFile(dateien_zip, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for p in sorted(DATEIEN.glob("*.xlsx")):
            zf.write(p, arcname=f"dateien/{p.name}")
        zf.write(hub, arcname="Cockpit-Hub-SharePoint.xlsx")
        zf.writestr(
            "README.txt",
            "Alle Excel-Dateien mit eingebackenen Werten (kein #WERT! durch Verknüpfungen).\n"
            "Hub als Cockpit.xlsx hochladen; restliche Dateien nach TVF-Controlling-Cockpit/.\n",
        )

    # Mirror packages
    for name in (
        "TVF_Cockpit_Download.zip",
        "TVF_Cockpit_Dateien.zip",
        "Cockpit.xlsx",
        "Cockpit-Hub-SharePoint.xlsx",
    ):
        src = ROOT / name
        if not src.exists():
            continue
        for dest_root in (ROOT.parent, STORE_DOCS, STORE_DOCS / "workbook-split"):
            shutil.copy2(src, dest_root / name)


def update_docs() -> None:
    note = ROOT / "WERTE-EINGEBACKEN.md"
    note.write_text(
        """# Werte eingebacken (SharePoint)

Excel Online zeigt bei Querverweisen oft `#WERT!` und rechnet Formeln nicht wie Desktop-Excel.

**Lösung:** `scripts/bake_all_sharepoint.py` schreibt die zuletzt berechneten Werte aus
`TVF_Controlling_Cockpit_Import.xlsx` als feste Zahlen in Hub + `dateien/*.xlsx`.

## Neu hochladen

1. `Cockpit-Hub-SharePoint.xlsx` → als `Cockpit.xlsx` ersetzen
2. Alle Dateien aus `dateien/` (oder `TVF_Cockpit_Dateien.zip`) ersetzen – besonders `Stunden.xlsx`, `KLR.xlsx`, `Steuerung.xlsx`, `GuV.xlsx`, …
3. SharePoint-Seite **veröffentlichen**
4. Hard-Reload: `Cmd + Shift + R`

## Erneut backen

```bash
python3 scripts/bake_all_sharepoint.py
```
""",
        encoding="utf-8",
    )
    # Patch ZUERST-LESEN
    zuerst = ROOT / "ZUERST-LESEN.md"
    if zuerst.exists():
        text = zuerst.read_text(encoding="utf-8")
        marker = "## Werte eingebacken"
        block = (
            "\n## Werte eingebacken\n\n"
            "Hub + `dateien/*.xlsx` enthalten **feste Zahlen** (kein Excel-Online-Neurechnen). "
            "Details: `WERTE-EINGEBACKEN.md`. Bei `#WERT!` Dateien neu hochladen und Hard-Reload.\n"
        )
        if marker not in text:
            zuerst.write_text(text.rstrip() + "\n" + block, encoding="utf-8")


def main() -> None:
    import_path = find_import()
    print(f"Import: {import_path}")
    print("Lade Cache (data_only) …")
    wb_val = load_workbook(import_path, data_only=True)
    print("Lade Formeln …")
    wb_src = load_workbook(import_path, data_only=False)

    print("Backe Hub …")
    rebuild_hub(wb_val, wb_src)

    print("Backe dateien/ …")
    bake_dateien(wb_val, wb_src)

    print("Pakete …")
    refresh_packages()
    update_docs()

    # quick verify
    wb = load_workbook(ROOT / "Cockpit.xlsx", data_only=False)
    c = wb["Cockpit"]
    print(
        "Check Cockpit:",
        "C6=", c["C6"].value,
        "G6=", c["G6"].value,
        "K6=", c["K6"].value,
        "O6=", c["O6"].value,
    )
    st = load_workbook(DATEIEN / "Stunden.xlsx", data_only=False)
    print("Check Stunden D7=", st["Stunden"]["D7"].value, "D8=", st["Stunden"]["D8"].value)
    print("fertig.")


if __name__ == "__main__":
    main()
