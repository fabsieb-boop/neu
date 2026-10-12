#!/usr/bin/env python3
"""Durchsuchbare Projektakten: eine Excel mit AutoFilter + Auswahl.

Problem: 168 Einzeldateien ohne Filter – man kann nichts suchen/anziehen.
Lösung: Projektakten_Suche.xlsx / Projektakte.xlsx mit
  - Alle_Projekte (AutoFilter, Freeze)
  - Auswahl (Projekt-Nr. gelb → Kennzahlen)
  - Rechnungen / Forecast / Stunden als filterbare Listen
"""

from __future__ import annotations

import re
import shutil
import zipfile
import warnings
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

warnings.filterwarnings("ignore")

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "projektakten"
DATEIEN = ROOT / "dateien"
ZIP_SRC = ROOT / "TVF_Projektakten.zip"
D_FINANZEN = DATEIEN / "D Finanzen.xlsx"
D_PROJEKTE = DATEIEN / "D Projekte.xlsx"
STORE = Path("/cursor/stores/bc-d9b5c533-1abd-48a4-a0b7-3919e592c358/docs")
ARTIFACTS = Path("/opt/cursor/artifacts")

HDR = Font(bold=True, color="FFFFFF", name="Calibri")
HDR_FILL = PatternFill("solid", fgColor="1F4E79")
TITLE = Font(bold=True, size=16, color="1F4E79", name="Calibri")
SUB = Font(size=10, color="555555", name="Calibri")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
MONEY = '#,##0.00'
THIN = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)


def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.font = HDR
        cell.fill = HDR_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = THIN


def autosize(ws, max_w=36):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        w = 11
        for cell in col[:80]:
            if cell.value is not None:
                w = min(max_w, max(w, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = w


def ensure_extracted():
    OUT.mkdir(parents=True, exist_ok=True)
    existing = list(OUT.glob("Projekt_*.xlsx"))
    if len(existing) >= 50:
        return
    if not ZIP_SRC.exists():
        raise SystemExit(f"ZIP fehlt: {ZIP_SRC}")
    print(f"Entpacke {ZIP_SRC.name} …")
    with zipfile.ZipFile(ZIP_SRC) as zf:
        for name in zf.namelist():
            if not name.startswith("projektakten/") or name.endswith("/"):
                continue
            target = ROOT / name
            target.parent.mkdir(parents=True, exist_ok=True)
            with zf.open(name) as src, open(target, "wb") as dst:
                shutil.copyfileobj(src, dst)


def read_block(ws, marker: str):
    start = None
    for r in range(1, (ws.max_row or 1) + 1):
        if ws.cell(r, 1).value == marker:
            start = r
            break
    if start is None:
        return [], []
    headers = []
    for c in range(1, 40):
        v = ws.cell(start + 1, c).value
        if v is None and c > 2:
            break
        headers.append(v)
    while headers and headers[-1] is None:
        headers.pop()
    rows = []
    for r in range(start + 2, (ws.max_row or 1) + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and a.startswith("v_"):
            break
        if all(ws.cell(r, c).value is None for c in range(1, len(headers) + 1)):
            break
        rows.append({headers[i]: ws.cell(r, i + 1).value for i in range(len(headers))})
    return headers, rows


def load_sources():
    print("Lese D Finanzen / D Projekte …")
    df = load_workbook(D_FINANZEN, data_only=True)["D Finanzen"]
    _, fc_auftrag = read_block(df, "v_forecast_auftrag")
    dp = load_workbook(D_PROJEKTE, data_only=True)["D Projekte"]
    _, projekt_db = read_block(dp, "v_projekt_db")
    _, kv_sap = read_block(dp, "v_kv_sap")
    _, stunden = read_block(dp, "v_stunden_projekt")
    _, bestand = read_block(dp, "v_bestand_ohne_rechnung")
    return {
        "fc": {str(r.get("auftrag")): r for r in fc_auftrag if r.get("auftrag")},
        "db": {str(r.get("projekt_nr")): r for r in projekt_db if r.get("projekt_nr")},
        "kv": {str(r.get("auftrag")): r for r in kv_sap if r.get("auftrag")},
        "stunden": stunden,
        "bestand": {str(r.get("auftrag")): r for r in bestand if r.get("auftrag")},
    }


def collect_from_files():
    """Pull Rechnungen (+ meta) from each Projekt_*.xlsx."""
    rechnungen = []
    forecast_m = []
    index = []
    files = sorted(OUT.glob("Projekt_*.xlsx"))
    print(f"Lese {len(files)} Einzelakten …")
    for i, path in enumerate(files, 1):
        wb = load_workbook(path, data_only=True)
        # meta from Uebersicht
        ue = wb["Uebersicht"] if "Uebersicht" in wb.sheetnames else wb.active
        meta = {}
        for r in range(4, 20):
            k, v = ue.cell(r, 1).value, ue.cell(r, 2).value
            if k:
                meta[str(k)] = v
        nr = str(meta.get("Projekt-Nr.") or "")
        name = str(meta.get("Projektname") or path.stem)
        if not nr:
            m = re.match(r"Projekt_([PW]\d+)_", path.name)
            nr = m.group(1) if m else path.stem
        index.append({
            "nr": nr,
            "name": name,
            "region": meta.get("Region"),
            "status": meta.get("Status"),
            "art": meta.get("Projektart"),
            "datei": path.name,
        })
        if "Rechnungen" in wb.sheetnames:
            ws = wb["Rechnungen"]
            # find header row
            hdr_row = None
            for r in range(1, 8):
                vals = [ws.cell(r, c).value for c in range(1, 10)]
                if "Belegnummer" in vals or "buchungsdatum" in [str(x).lower() if x else "" for x in vals] or (
                    vals and str(vals[0]).lower() in ("datum", "buchungsdatum", "belegart")
                ):
                    hdr_row = r
                    break
                if vals and any(x in ("Belegart", "Kostenart", "Wert", "Periode") for x in vals if x):
                    hdr_row = r
                    break
            if hdr_row:
                headers = [ws.cell(hdr_row, c).value for c in range(1, 12)]
                for r in range(hdr_row + 1, min(ws.max_row, hdr_row + 500) + 1):
                    rowvals = [ws.cell(r, c).value for c in range(1, 12)]
                    if all(v is None for v in rowvals):
                        continue
                    rec = {"auftrag": nr, "projektname": name}
                    for h, v in zip(headers, rowvals):
                        if h:
                            rec[str(h)] = v
                    rechnungen.append(rec)
        if "Forecast_Monat" in wb.sheetnames:
            ws = wb["Forecast_Monat"]
            for r in range(1, min(ws.max_row, 30) + 1):
                # look for month numbers in col A/C
                for c0 in (1, 3):
                    m = ws.cell(r, c0).value
                    if isinstance(m, int) and 1 <= m <= 12:
                        forecast_m.append({
                            "auftrag": nr,
                            "projektname": name,
                            "monat": m,
                            "ist": ws.cell(r, c0 + 1).value,
                        })
                        break
        wb.close()
        if i % 40 == 0:
            print(f"  … {i}")
    return index, rechnungen, forecast_m


def add_table(ws, ref: str, name: str):
    # Excel table names: no spaces
    safe = re.sub(r"[^A-Za-z0-9_]", "_", name)[:40]
    tab = Table(displayName=safe, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False,
        showLastColumn=False, showRowStripes=True,
    )
    ws.add_table(tab)


def build_workbook(sources, index, rechnungen, forecast_m) -> Workbook:
    wb = Workbook()

    # ---------- Suche / Anleitung ----------
    ws = wb.active
    ws.title = "Suche"
    ws["B2"] = "Projektakten – Suche & Filter"
    ws["B2"].font = TITLE
    ws["B3"] = (
        "Alles in EINER Datei. Blatt „Alle_Projekte“: AutoFilter an (Dropdown in Kopfzeile) – "
        "nach Region, Status, Name, Auftrag filtern. "
        "Blatt „Auswahl“: Projekt-Nr. in die gelbe Zelle → Kennzahlen erscheinen."
    )
    ws["B3"].font = SUB
    ws.merge_cells("B3:H3")
    ws["B5"] = "So geht’s"
    ws["B5"].font = Font(bold=True, size=12, color="1F4E79")
    steps = [
        "1. Blatt Alle_Projekte öffnen",
        "2. In der Kopfzeile auf den Pfeil klicken (Filter)",
        "3. Region / Status wählen oder im Namensfilter tippen",
        "4. Gesuchtes Projekt-Nr. kopieren → Blatt Auswahl, gelbe Zelle",
        "5. Rechnungen / Forecast_Liste: ebenfalls Filter auf Spalte Auftrag",
    ]
    for i, s in enumerate(steps, 7):
        ws.cell(i, 2, s)
    ws["B14"] = "Einzelakten (optional)"
    ws["B14"].font = Font(bold=True, color="1F4E79")
    ws["B15"] = (
        "Die bisherigen Dateien Projekt_….xlsx liegen weiter im Ordner projektakten/ "
        "für den Deep-Dive. Für den Alltag reicht diese Such-Datei."
    )
    ws["B15"].font = SUB
    autosize(ws, 80)

    # ---------- Alle_Projekte ----------
    ws = wb.create_sheet("Alle_Projekte")
    ws["A1"] = "Alle Projekte – filtern / suchen"
    ws["A1"].font = TITLE
    ws["A2"] = "Filter in Zeile 4 nutzen. Spalte „Suchtext“ enthält Nr+Name+Region für Textsuche im Filter."
    ws["A2"].font = SUB

    headers = [
        "Projekt-Nr.", "Name", "Region", "Status", "Art",
        "Leistung", "Var. Kosten", "DB4",
        "Umsatz Plan (FC)", "Umsatz Ist (Rechnung)", "Offen FC−Ist",
        "Kosten Plan", "Kosten Ist",
        "Umsatz KV", "Umsatz SAP", "Δ Umsatz KV−SAP",
        "Auftraggeber", "Phase", "Erfüllung %",
        "Tage ohne Rechnung", "Datei", "Suchtext",
    ]
    hr = 4
    for j, h in enumerate(headers, 1):
        ws.cell(hr, j, h)
    style_header(ws, hr, len(headers))

    # merge index with sources
    by_nr = {x["nr"]: x for x in index}
    all_nrs = sorted(set(by_nr) | set(sources["db"]) | set(sources["fc"]) | set(sources["kv"]))

    data_start = hr + 1
    for i, nr in enumerate(all_nrs):
        row_i = data_start + i
        idx = by_nr.get(nr, {"nr": nr, "name": nr, "region": None, "status": None, "art": None, "datei": ""})
        db = sources["db"].get(nr, {})
        fc = sources["fc"].get(nr, {})
        kv = sources["kv"].get(nr, {})
        best = sources["bestand"].get(nr, {})
        name = idx.get("name") or db.get("projektname") or kv.get("projekttitel") or nr
        region = idx.get("region") or db.get("region") or kv.get("region")
        status = idx.get("status") or db.get("status") or kv.get("status")
        up = fc.get("umsatz_plan") or 0
        ui = fc.get("umsatz_ist") or 0
        try:
            upf, uif = float(up or 0), float(ui or 0)
        except (TypeError, ValueError):
            upf, uif = 0.0, 0.0
        ukv = kv.get("umsatz_kv")
        usap = kv.get("umsatz_sap")
        try:
            delta_kv = (float(ukv or 0) - float(usap or 0)) if ukv is not None or usap is not None else None
        except (TypeError, ValueError):
            delta_kv = None
        such = f"{nr} {name} {region or ''} {status or ''} {kv.get('auftraggeber') or ''}".lower()
        vals = [
            nr, name, region, status, idx.get("art") or db.get("projektart") or kv.get("phase"),
            db.get("leistung"), db.get("variable_kosten"), db.get("db4"),
            upf, uif, upf - uif,
            fc.get("kosten_plan") or 0, fc.get("kosten_ist") or 0,
            ukv, usap, delta_kv,
            kv.get("auftraggeber"), kv.get("phase") or kv.get("status"), kv.get("erfuellung_prozent"),
            best.get("tage_ohne_rechnung"), idx.get("datei") or "", such,
        ]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(row_i, j, v)
            cell.border = THIN
            if j in (6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16) and isinstance(v, (int, float)):
                cell.number_format = MONEY
        # hyperlink to file if present
        if idx.get("datei"):
            cell = ws.cell(row_i, 21)
            cell.hyperlink = idx["datei"]
            cell.style = "Hyperlink"

    last = data_start + len(all_nrs) - 1
    if all_nrs:
        add_table(ws, f"A{hr}:{get_column_letter(len(headers))}{last}", "AlleProjekte")
        ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(headers))}{last}"
        ws.freeze_panes = f"A{data_start}"
    autosize(ws)
    # hide Suchtext helper? keep visible for filter search
    ws.column_dimensions["V"].width = 28

    # ---------- Auswahl ----------
    ws = wb.create_sheet("Auswahl", 1)
    ws["B2"] = "Projekt anzeigen"
    ws["B2"].font = TITLE
    ws["B3"] = "Gelbe Zelle: Projekt-Nr. wählen (Dropdown oder eintippen). Kennzahlen kommen aus Alle_Projekte."
    ws["B3"].font = SUB
    ws["B5"] = "Projekt-Nr."
    ws["B5"].font = Font(bold=True)
    ws["C5"] = all_nrs[0] if all_nrs else ""
    ws["C5"].fill = YELLOW
    ws["C5"].border = THIN
    ws["C5"].font = Font(bold=True, size=14, color="1F4E79")

    # dropdown
    if all_nrs:
        # list on hidden area or use formula list – put nrs on Alle_Projekte col A
        dv = DataValidation(
            type="list",
            formula1=f"=Alle_Projekte!$A${data_start}:$A${last}",
            allow_blank=True,
        )
        dv.error = "Projekt-Nr. aus der Liste wählen"
        dv.errorTitle = "Unbekannt"
        ws.add_data_validation(dv)
        dv.add(ws["C5"])

    labels = [
        (7, "Name", 2),
        (8, "Region", 3),
        (9, "Status", 4),
        (10, "Art / Phase", 5),
        (11, "Leistung", 6),
        (12, "Var. Kosten", 7),
        (13, "DB4", 8),
        (14, "Umsatz Plan (FC)", 9),
        (15, "Umsatz Ist (Rechnung)", 10),
        (16, "Offen FC−Ist", 11),
        (17, "Kosten Plan", 12),
        (18, "Kosten Ist", 13),
        (19, "Umsatz KV", 14),
        (20, "Umsatz SAP", 15),
        (21, "Auftraggeber", 17),
        (22, "Tage ohne Rechnung", 20),
        (23, "Einzeldatei", 21),
    ]
    ws["B6"] = "Kennzahl"
    ws["C6"] = "Wert"
    style_header(ws, 6, 2)
    for row, label, col in labels:
        ws.cell(row, 2, label).border = THIN
        # XLOOKUP / INDEX MATCH – Excel Online supports XLOOKUP
        col_letter = get_column_letter(col)
        formula = (
            f'=IFERROR(XLOOKUP($C$5,Alle_Projekte!$A${data_start}:$A${last},'
            f'Alle_Projekte!${col_letter}${data_start}:${col_letter}${last}),"")'
        )
        cell = ws.cell(row, 3, formula)
        cell.border = THIN
        if col in (6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16):
            cell.number_format = MONEY

    ws["B25"] = "Rechnungen dieses Projekts → Blatt Rechnungen, Filter auf Auftrag = gelbe Nr."
    ws["B25"].font = SUB
    ws["B26"] = "Oder AutoFilter in Alle_Projekte auf genau diese Nr. setzen."
    ws["B26"].font = SUB
    autosize(ws, 40)
    ws.column_dimensions["C"].width = 28

    # ---------- Rechnungen ----------
    ws = wb.create_sheet("Rechnungen")
    ws["A1"] = "Rechnungen / Erlöse aller Projekte"
    ws["A1"].font = TITLE
    ws["A2"] = "AutoFilter: Spalte Auftrag = gewünschte Projekt-Nr. (aus Auswahl kopieren)."
    ws["A2"].font = SUB
    # unify headers
    rh = ["Auftrag", "Projektname", "Buchungsdatum", "Periode", "Jahr", "Belegart", "Belegnummer", "Kostenart", "Wert"]
    # map common keys from collected dicts
    key_map = {
        "Buchungsdatum": ("Buchungsdatum", "buchungsdatum", "Datum"),
        "Periode": ("Periode", "periode"),
        "Jahr": ("Jahr", "jahr"),
        "Belegart": ("Belegart", "belegart"),
        "Belegnummer": ("Belegnummer", "belegnummer"),
        "Kostenart": ("Kostenart", "kostenart"),
        "Wert": ("Wert", "wert"),
    }
    hr = 4
    for j, h in enumerate(rh, 1):
        ws.cell(hr, j, h)
    style_header(ws, hr, len(rh))
    # cap for file size
    MAX_R = 8000
    for i, rec in enumerate(rechnungen[:MAX_R]):
        row = hr + 1 + i
        def pick(keys):
            for k in keys:
                if k in rec and rec[k] is not None:
                    return rec[k]
            return None
        vals = [
            rec.get("auftrag"),
            rec.get("projektname"),
            pick(key_map["Buchungsdatum"]),
            pick(key_map["Periode"]),
            pick(key_map["Jahr"]),
            pick(key_map["Belegart"]),
            pick(key_map["Belegnummer"]),
            pick(key_map["Kostenart"]),
            pick(key_map["Wert"]),
        ]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(row, j, v)
            cell.border = THIN
            if j == 9 and isinstance(v, (int, float)):
                cell.number_format = MONEY
    last_r = hr + min(len(rechnungen), MAX_R)
    if rechnungen:
        add_table(ws, f"A{hr}:I{last_r}", "RechnungenAlle")
        ws.auto_filter.ref = f"A{hr}:I{last_r}"
        ws.freeze_panes = f"A{hr+1}"
    if len(rechnungen) > MAX_R:
        ws.cell(last_r + 2, 1, f"Hinweis: auf {MAX_R} von {len(rechnungen)} Rechnungszeilen begrenzt.")
    autosize(ws)

    # ---------- Forecast_Liste ----------
    ws = wb.create_sheet("Forecast_Liste")
    ws["A1"] = "Forecast je Auftrag (filterbar)"
    ws["A1"].font = TITLE
    ws["A2"] = "Dieselbe Sicht wie Forecast je Auftrag.xlsx – hier mit AutoFilter."
    ws["A2"].font = SUB
    fh = [
        "Auftrag", "Name", "Region", "Umsatz Plan", "Umsatz Ist", "Offen",
        "Umsatz Restplan", "Umsatz erwartet", "Kosten Plan", "Kosten Ist", "Jahr",
    ]
    hr = 4
    for j, h in enumerate(fh, 1):
        ws.cell(hr, j, h)
    style_header(ws, hr, len(fh))
    fc_rows = list(sources["fc"].values())
    fc_rows.sort(key=lambda r: str(r.get("auftrag") or ""))
    for i, r in enumerate(fc_rows):
        row = hr + 1 + i
        up = float(r.get("umsatz_plan") or 0)
        ui = float(r.get("umsatz_ist") or 0)
        vals = [
            r.get("auftrag"), r.get("auftragname"), r.get("region"),
            up, ui, up - ui,
            r.get("umsatz_restplan") or 0, r.get("umsatz_erwartet") or 0,
            r.get("kosten_plan") or 0, r.get("kosten_ist") or 0,
            r.get("jahr") or 2026,
        ]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(row, j, v)
            cell.border = THIN
            if j in range(4, 11) and isinstance(v, (int, float)):
                cell.number_format = MONEY
    last_f = hr + len(fc_rows)
    if fc_rows:
        add_table(ws, f"A{hr}:{get_column_letter(len(fh))}{last_f}", "ForecastListe")
        ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(fh))}{last_f}"
        ws.freeze_panes = f"A{hr+1}"
    autosize(ws)

    # ---------- Stunden ----------
    ws = wb.create_sheet("Stunden")
    ws["A1"] = "Projektstunden (filterbar)"
    ws["A1"].font = TITLE
    sh = ["Jahr", "Monat", "Projekt-Nr.", "Region", "Name", "Stunden", "Mitarbeiter", "Tage"]
    hr = 4
    for j, h in enumerate(sh, 1):
        ws.cell(hr, j, h)
    style_header(ws, hr, len(sh))
    for i, r in enumerate(sources["stunden"]):
        row = hr + 1 + i
        vals = [
            r.get("jahr"), r.get("monat"), r.get("projekt_nr"), r.get("region"),
            r.get("projektname"), r.get("stunden"), r.get("mitarbeiter"), r.get("tage"),
        ]
        for j, v in enumerate(vals, 1):
            ws.cell(row, j, v).border = THIN
    last_s = hr + len(sources["stunden"])
    if sources["stunden"]:
        add_table(ws, f"A{hr}:H{last_s}", "StundenAlle")
        ws.auto_filter.ref = f"A{hr}:H{last_s}"
        ws.freeze_panes = f"A{hr+1}"
    autosize(ws)

    return wb


def fix_uebersicht_autofilter():
    path = OUT / "00_Projektakten_Uebersicht.xlsx"
    if not path.exists():
        return
    wb = load_workbook(path)
    ws = wb.active
    # find header row
    hr = 4
    max_r = ws.max_row
    max_c = ws.max_column
    ws.auto_filter.ref = f"A{hr}:{get_column_letter(max_c)}{max_r}"
    ws.freeze_panes = f"A{hr+1}"
    # add Suchtext column
    such_col = max_c + 1
    ws.cell(hr, such_col, "Suchtext")
    style_header(ws, hr, such_col)
    for r in range(hr + 1, max_r + 1):
        parts = [ws.cell(r, c).value for c in range(1, min(5, max_c) + 1)]
        ws.cell(r, such_col, " ".join(str(p) for p in parts if p).lower())
    ws.auto_filter.ref = f"A{hr}:{get_column_letter(such_col)}{max_r}"
    # hyperlinks on Datei column if present
    datei_col = None
    for c in range(1, max_c + 1):
        if ws.cell(hr, c).value in ("Datei", "datei"):
            datei_col = c
            break
    if datei_col:
        for r in range(hr + 1, max_r + 1):
            fname = ws.cell(r, datei_col).value
            if fname and (OUT / str(fname)).exists():
                ws.cell(r, datei_col).hyperlink = str(fname)
                ws.cell(r, datei_col).style = "Hyperlink"
    wb.save(path)


def pack():
    search = OUT / "Projektakten_Suche.xlsx"
    # also as dateien/Projektakte.xlsx (SharePoint main entry)
    shutil.copy2(search, DATEIEN / "Projektakte.xlsx")
    shutil.copy2(search, DATEIEN / "Projektakten_Suche.xlsx")

    z = ROOT / "TVF_Projektakten.zip"
    with zipfile.ZipFile(z, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.write(search, "projektakten/Projektakten_Suche.xlsx")
        zf.write(search, "Projektakten_Suche.xlsx")
        for f in sorted(OUT.iterdir()):
            if f.is_file() and f.name != "Projektakten_Suche.xlsx":
                zf.write(f, f"projektakten/{f.name}")
        zf.writestr(
            "ZUERST-LESEN.txt",
            "WICHTIG: Öffne zuerst Projektakten_Suche.xlsx\n"
            "→ Blatt Alle_Projekte: Filter in der Kopfzeile\n"
            "→ Blatt Auswahl: Projekt-Nr. in gelbe Zelle\n"
            "Einzeldateien Projekt_….xlsx nur für Detailtiefe.\n",
        )

    small = ROOT / "TVF_Projektakten_Suche.zip"
    with zipfile.ZipFile(small, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.write(search, "Projektakten_Suche.xlsx")
        zf.writestr(
            "README.txt",
            "Eine Datei zum Suchen/Filtern/Anziehen von Projekten.\n"
            "Nach TVF-Controlling-Cockpit/ hochladen (als Projektakte.xlsx ersetzen).\n",
        )

    for name in (
        "TVF_Projektakten.zip",
        "TVF_Projektakten_Suche.zip",
        "Projektakten_Suche.xlsx",
    ):
        src = ROOT / name if not name.endswith("Suche.xlsx") else OUT / "Projektakten_Suche.xlsx"
        if name == "Projektakten_Suche.xlsx":
            src = OUT / name
        if not src.exists():
            continue
        for dest in (ROOT.parent, STORE, STORE / "workbook-split", ARTIFACTS):
            dest.mkdir(parents=True, exist_ok=True)
            shutil.copy2(src, dest / (name if name != "Projektakten_Suche.xlsx" else name))
    shutil.copy2(DATEIEN / "Projektakte.xlsx", ARTIFACTS / "Projektakte.xlsx")
    shutil.copy2(DATEIEN / "Projektakte.xlsx", STORE / "workbook-split" / "dateien" / "Projektakte.xlsx")


def update_readme():
    (OUT / "README.md").write_text(
        """# Projektakten

## Zuerst öffnen: `Projektakten_Suche.xlsx`

Eine Datei – suchen, filtern, Projekt anziehen:

| Blatt | Nutzen |
|-------|--------|
| **Suche** | Kurzanleitung |
| **Alle_Projekte** | AutoFilter (Region, Status, Name, Suchtext) |
| **Auswahl** | Projekt-Nr. in gelbe Zelle → Kennzahlen |
| **Rechnungen** | Alle Rechnungen, Filter auf Auftrag |
| **Forecast_Liste** | Forecast Plan/Ist je Auftrag |
| **Stunden** | Projektstunden filterbar |

Die Einzeldateien `Projekt_….xlsx` bleiben für den Detail-Deep-Dive.

```bash
python3 scripts/build_projektakten_suche.py
```

Downloads: `TVF_Projektakten_Suche.zip` (nur Suche) oder `TVF_Projektakten.zip` (alles).
""",
        encoding="utf-8",
    )
    (ROOT / "PROJEKTAKTEN-SUCHE.md").write_text(
        (OUT / "README.md").read_text(encoding="utf-8"), encoding="utf-8"
    )


def main():
    ensure_extracted()
    sources = load_sources()
    index, rechnungen, forecast_m = collect_from_files()
    print(f"Index {len(index)}, Rechnungen {len(rechnungen)}")
    wb = build_workbook(sources, index, rechnungen, forecast_m)
    out = OUT / "Projektakten_Suche.xlsx"
    wb.save(out)
    print("geschrieben", out, out.stat().st_size)
    fix_uebersicht_autofilter()
    update_readme()
    pack()
    # verify
    v = load_workbook(out)
    print("Blätter:", v.sheetnames)
    print("Alle_Projekte Filter:", v["Alle_Projekte"].auto_filter.ref)
    print("fertig.")


if __name__ == "__main__":
    main()
